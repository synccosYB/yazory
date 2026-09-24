import json
import os
import re
import csv
import hashlib
import hmac
from decimal import Decimal, InvalidOperation
from email.utils import getaddresses
from io import BytesIO, StringIO

import app_original as _app

from flask import Response, current_app, has_request_context, jsonify, session
from sqlalchemy import Index, UniqueConstraint, case, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload, selectinload
from sqlalchemy.orm.exc import StaleDataError
from werkzeug.exceptions import Forbidden
from twilio_service import (account_overview, create_messaging_service,
                            deliver_message, find_messaging_service_for_number,
                            message_status, normalize_phone,
                            validate_webhook_signature)


def _public_url(endpoint, **values):
    base = current_app.config.get('APP_BASE_URL', '').rstrip('/')
    return (base + _app.url_for(endpoint, **values) if base else
            _app.url_for(endpoint, _external=True, **values))


class SupporterProfile(_app.db.Model):
    """A person in the shared directory, before they are assigned to a case."""
    __tablename__ = 'supporter_profile'
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    person_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('supporter_person.id'), nullable=True,
        unique=True, index=True)
    name = _app.db.Column(_app.db.String(160), nullable=False)
    phone = _app.db.Column(_app.db.String(80), nullable=False)
    normalized_phone = _app.db.Column(
        _app.db.String(20), nullable=False, unique=True, index=True)
    email = _app.db.Column(_app.db.String(254), nullable=False, default='')
    created_at = _app.db.Column(
        _app.db.DateTime, nullable=False,
        default=lambda: _app.datetime.now(_app.timezone.utc).replace(tzinfo=None))


class PersonRelationship(_app.db.Model):
    """A direct connection between two canonical people, independent of a case."""
    __tablename__ = 'person_relationship'
    __table_args__ = (
        UniqueConstraint('person_one_id', 'person_two_id',
                         name='uq_person_relationship_pair'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    person_one_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('supporter_person.id'),
        nullable=False, index=True)
    person_two_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('supporter_person.id'),
        nullable=False, index=True)
    relationship = _app.db.Column(_app.db.String(80), nullable=False)
    notes = _app.db.Column(_app.db.String(500), nullable=False, default='')


PERSON_RELATIONSHIPS = (
    'Brothers', 'Sisters', 'Brother and sister', 'Spouses',
    'Parent and child', 'In-laws', 'Cousins', 'Friends', 'Other')


def normalized_profile_phone(value):
    """Return the canonical phone identity used by imports and case contacts."""
    digits = re.sub(r'\D', '', value or '')
    if len(digits) == 11 and digits.startswith('1'):
        digits = digits[1:]
    return digits if 7 <= len(digits) <= 15 else ''


def imported_contact_rows(upload):
    """Read a small CSV or XLSX upload and yield normalized field dictionaries."""
    filename = (upload.filename or '').lower()
    raw = upload.read()
    if filename.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency is deployed
            raise ValueError('Excel imports are unavailable on this server.') from exc
        sheet = load_workbook(BytesIO(raw), read_only=True, data_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value or '').strip() for value in values[0]]
        source_rows = (dict(zip(headers, row)) for row in values[1:])
    elif filename.endswith('.csv'):
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('Save the CSV as UTF-8 and upload it again.') from exc
        source_rows = csv.DictReader(StringIO(content))
    else:
        raise ValueError('Upload a CSV or Excel (.xlsx) file.')

    # Ordered aliases support both a simple three-column file and exports with
    # separate English/Yiddish names and several phone/email columns.  The
    # first populated phone is the profile's one unique identity.
    aliases = {
        'name': ('yiddish/hebrew name', 'english name', 'name', 'full name',
                 'fullname', 'supporter', 'contact', 'original full name'),
        'phone': ('phone 1', 'phone1', 'phone', 'phone number', 'telephone',
                  'cell', 'cell phone', 'mobile', 'phone 2', 'phone2',
                  'phone 3', 'phone3'),
        'email': ('email 1', 'email1', 'email', 'email address', 'e-mail',
                  'email 2', 'email2'),
    }
    rows = []
    for number, source in enumerate(source_rows, start=2):
        cleaned = {
            re.sub(r'[_-]+', ' ', str(key or '').strip().lower()):
            str(value or '').strip() for key, value in source.items()
        }
        row = {'row': number}
        for field, choices in aliases.items():
            row[field] = next((cleaned[key] for key in choices if cleaned.get(key)), '')
        rows.append(row)
    return rows

if 'Shul friend' not in _app.RELATIONSHIPS:
    insert_at = _app.RELATIONSHIPS.index('Friend') if 'Friend' in _app.RELATIONSHIPS else len(_app.RELATIONSHIPS)
    _app.RELATIONSHIPS.insert(insert_at, 'Shul friend')

UNCLE_RELATIONSHIPS = [
    "Applicant’s uncle — father’s brother",
    "Applicant’s uncle — mother’s brother",
    "Applicant’s uncle — father’s sister’s husband",
    "Applicant’s uncle — mother’s sister’s husband",
]
for _relationship in reversed(UNCLE_RELATIONSHIPS):
    if _relationship not in _app.RELATIONSHIPS:
        _app.RELATIONSHIPS.insert(
            _app.RELATIONSHIPS.index('First cousin'), _relationship)


class ShulRabbi(_app.db.Model):
    """Exactly one rabbi assignment for a shared shul record."""
    __tablename__ = 'shul_rabbi'
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), primary_key=True)
    rabbi_name = _app.db.Column(_app.db.String(160), nullable=False, default='')
    rabbi_phone = _app.db.Column(_app.db.String(80), nullable=False, default='')
    institution = _app.db.relationship('Institution')


class ShulRabbiPhone(_app.db.Model):
    """All phone numbers that belong to the rabbi of a shul."""
    __tablename__ = 'shul_rabbi_phone'
    __table_args__ = (
        UniqueConstraint('institution_id', 'phone', name='uq_shul_rabbi_phone'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), nullable=False, index=True)
    phone = _app.db.Column(_app.db.String(80), nullable=False)
    institution = _app.db.relationship('Institution')


class ShulRabbiAssistant(_app.db.Model):
    """An assistant/gabbai of the rabbi, distinct from a gabbai of the shul."""
    __tablename__ = 'shul_rabbi_assistant'
    __table_args__ = (
        UniqueConstraint('institution_id', 'name', name='uq_shul_rabbi_assistant'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), nullable=False, index=True)
    name = _app.db.Column(_app.db.String(160), nullable=False)
    institution = _app.db.relationship('Institution')


class ShulRabbiAssistantPhone(_app.db.Model):
    """Multiple phone numbers for one rabbi assistant."""
    __tablename__ = 'shul_rabbi_assistant_phone'
    __table_args__ = (
        UniqueConstraint('assistant_id', 'phone', name='uq_shul_rabbi_assistant_phone'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    assistant_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('shul_rabbi_assistant.id'), nullable=False, index=True)
    phone = _app.db.Column(_app.db.String(80), nullable=False)
    assistant = _app.db.relationship('ShulRabbiAssistant')


class RabbiPerson(_app.db.Model):
    """One canonical rabbi identity, reusable by any number of shuls."""
    __tablename__ = 'rabbi_person'
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    name = _app.db.Column(_app.db.String(160), nullable=False)
    normalized_name = _app.db.Column(_app.db.String(160), nullable=False, unique=True, index=True)
    phone = _app.db.Column(_app.db.String(80), nullable=False, default='')


class ShulRabbiAssociation(_app.db.Model):
    """The many-to-many shul/rabbi directory, including its single primary."""
    __tablename__ = 'shul_rabbi_association'
    __table_args__ = (
        UniqueConstraint('institution_id', 'rabbi_person_id',
                         name='uq_shul_rabbi_association'),
        Index('uq_shul_rabbi_primary', 'institution_id', unique=True,
              sqlite_where=text('is_primary = 1'),
              postgresql_where=text('is_primary = true')),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), nullable=False, index=True)
    rabbi_person_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('rabbi_person.id'), nullable=False, index=True)
    is_primary = _app.db.Column(_app.db.Boolean, nullable=False, default=False)
    institution = _app.db.relationship(
        'Institution', backref=_app.db.backref(
            'rabbi_associations', cascade='all, delete-orphan',
            order_by='ShulRabbiAssociation.id'))
    rabbi_person = _app.db.relationship('RabbiPerson')


class FamilyRabbiPreference(_app.db.Model):
    """Canonical family choice; legacy Family columns remain display snapshots."""
    __tablename__ = 'family_rabbi_preference'
    family_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('family.id'), primary_key=True)
    rabbi_person_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('rabbi_person.id'), nullable=True, index=True)
    overridden = _app.db.Column(_app.db.Boolean, nullable=False, default=False)
    rabbi_person = _app.db.relationship('RabbiPerson')


class HelperPerson(_app.db.Model):
    """One reusable helper identity that may serve any number of shuls."""
    __tablename__ = 'helper_person'
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    name = _app.db.Column(_app.db.String(160), nullable=False)
    normalized_name = _app.db.Column(
        _app.db.String(160), nullable=False, unique=True, index=True)


class HelperPhone(_app.db.Model):
    """Phone numbers belong to the helper, rather than to one shul link."""
    __tablename__ = 'helper_phone'
    __table_args__ = (
        UniqueConstraint('helper_person_id', 'phone', name='uq_helper_phone'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    helper_person_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('helper_person.id'), nullable=False, index=True)
    phone = _app.db.Column(_app.db.String(80), nullable=False)
    helper_person = _app.db.relationship('HelperPerson')


class ShulHelperAssociation(_app.db.Model):
    """Many-to-many shul/helper link, with the helper's role at that shul."""
    __tablename__ = 'shul_helper_association'
    __table_args__ = (
        UniqueConstraint('institution_id', 'helper_person_id', 'role',
                         name='uq_shul_helper_association'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), nullable=False, index=True)
    helper_person_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('helper_person.id'), nullable=False, index=True)
    role = _app.db.Column(_app.db.String(40), nullable=False)
    institution = _app.db.relationship(
        'Institution', backref=_app.db.backref(
            'helper_associations', cascade='all, delete-orphan',
            order_by='ShulHelperAssociation.id'))
    helper_person = _app.db.relationship('HelperPerson')


class FamilyPhone(_app.db.Model):
    """Additional home phone numbers for an applicant household."""
    __tablename__ = 'family_phone'
    __table_args__ = (
        UniqueConstraint('family_id', 'phone', name='uq_family_phone'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    family_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('family.id'), nullable=False, index=True)
    phone = _app.db.Column(_app.db.String(80), nullable=False)


class FamilyRabbiConnection(_app.db.Model):
    """A family may have separate affiliated, weekday-shul, and Shabbos-shul rabbis."""
    __tablename__ = 'family_rabbi_connection'
    __table_args__ = (
        UniqueConstraint('family_id', 'role', name='uq_family_rabbi_connection_role'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    family_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('family.id'), nullable=False, index=True)
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), nullable=True, index=True)
    role = _app.db.Column(_app.db.String(30), nullable=False)
    rabbi_name = _app.db.Column(_app.db.String(160), nullable=False, default='')
    rabbi_phone = _app.db.Column(_app.db.String(80), nullable=False, default='')
    institution = _app.db.relationship('Institution')


class ShulGabbaiDirectory(_app.db.Model):
    """A shul can have multiple gabbaim, each stored once on the shul itself."""
    __tablename__ = 'shul_gabbai_directory'
    __table_args__ = (
        UniqueConstraint('institution_id', 'name', 'phone', name='uq_shul_gabbai_identity'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), nullable=False, index=True)
    name = _app.db.Column(_app.db.String(160), nullable=False)
    phone = _app.db.Column(_app.db.String(80), nullable=False, default='')
    institution = _app.db.relationship('Institution')


class ShulGabbaiPhone(_app.db.Model):
    """Multiple phone numbers for a shul gabbai."""
    __tablename__ = 'shul_gabbai_phone'
    __table_args__ = (
        UniqueConstraint('gabbai_id', 'phone', name='uq_shul_gabbai_phone'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    gabbai_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('shul_gabbai_directory.id'), nullable=False, index=True)
    phone = _app.db.Column(_app.db.String(80), nullable=False)
    gabbai = _app.db.relationship('ShulGabbaiDirectory')


class FamilyGabbaiConnection(_app.db.Model):
    """An applicant inherits all gabbaim from each selected shul."""
    __tablename__ = 'family_gabbai_connection'
    __table_args__ = (
        UniqueConstraint('family_id', 'role', 'gabbai_id', name='uq_family_gabbai_connection'),
    )
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    family_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('family.id'), nullable=False, index=True)
    institution_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('institution.id'), nullable=False, index=True)
    gabbai_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('shul_gabbai_directory.id'), nullable=False, index=True)
    role = _app.db.Column(_app.db.String(30), nullable=False)
    gabbai = _app.db.relationship('ShulGabbaiDirectory')
    institution = _app.db.relationship('Institution')


class StaffTask(_app.db.Model):
    """Assignable staff work; a parent task may contain any number of subtasks."""
    __tablename__ = 'staff_task'
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    parent_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('staff_task.id'), nullable=True, index=True)
    family_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('family.id'), nullable=True, index=True)
    source_contact_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('contact.id', ondelete='CASCADE'),
        nullable=True, unique=True, index=True)
    assigned_to = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('staff_user.id'), nullable=False, index=True)
    created_by = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('staff_user.id'), nullable=False, index=True)
    title = _app.db.Column(_app.db.String(240), nullable=False)
    description = _app.db.Column(_app.db.Text, nullable=False, default='')
    outcome = _app.db.Column(_app.db.Text, nullable=False, default='')
    status = _app.db.Column(_app.db.String(20), nullable=False, default='To do', index=True)
    priority = _app.db.Column(_app.db.String(20), nullable=False, default='Normal', index=True)
    due_date = _app.db.Column(_app.db.Date, nullable=True, index=True)
    created_at = _app.db.Column(
        _app.db.DateTime, nullable=False,
        default=lambda: _app.datetime.now(_app.timezone.utc).replace(tzinfo=None))
    completed_at = _app.db.Column(_app.db.DateTime, nullable=True)
    parent = _app.db.relationship(
        'StaffTask', remote_side=[id], backref=_app.db.backref(
            'subtasks', cascade='all, delete-orphan', order_by='StaffTask.id'))
    family = _app.db.relationship('Family')
    source_contact = _app.db.relationship('Contact', backref=_app.db.backref(
        'automatic_followup_task', uselist=False, cascade='all, delete-orphan',
        single_parent=True))
    assignee = _app.db.relationship('StaffUser', foreign_keys=[assigned_to])
    creator = _app.db.relationship('StaffUser', foreign_keys=[created_by])


class SupporterCommunication(_app.db.Model):
    """One chronological outreach, pledge, or receipt event for a supporter."""
    __tablename__ = 'supporter_communication'
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    contact_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('contact.id', ondelete='CASCADE'),
        nullable=False, index=True)
    family_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('family.id'), nullable=False, index=True)
    staff_user_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('staff_user.id'), nullable=True, index=True)
    email_message_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('email_message.id'), nullable=True, index=True)
    receipt_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('receipt.id'), nullable=True,
        unique=True, index=True)
    provider_message_id = _app.db.Column(_app.db.String(100), nullable=False, default='')
    delivery_error = _app.db.Column(_app.db.Text, nullable=False, default='')
    kind = _app.db.Column(_app.db.String(30), nullable=False, index=True)
    direction = _app.db.Column(_app.db.String(20), nullable=False, default='outbound')
    subject = _app.db.Column(_app.db.String(300), nullable=False, default='')
    body = _app.db.Column(_app.db.Text, nullable=False, default='')
    status = _app.db.Column(_app.db.String(20), nullable=False, default='completed', index=True)
    scheduled_for = _app.db.Column(_app.db.DateTime, nullable=True, index=True)
    created_at = _app.db.Column(
        _app.db.DateTime, nullable=False,
        default=lambda: _app.datetime.now(_app.timezone.utc).replace(tzinfo=None))
    completed_at = _app.db.Column(_app.db.DateTime, nullable=True)
    contact = _app.db.relationship('Contact')
    family = _app.db.relationship('Family')
    staff_user = _app.db.relationship('StaffUser')
    email_message = _app.db.relationship('EmailMessage')
    receipt = _app.db.relationship('Receipt')


class ResendWebhookEvent(_app.db.Model):
    """Deduplicate Resend's at-least-once delivery and manual replays."""
    __tablename__ = 'resend_webhook_event'
    event_id = _app.db.Column(_app.db.String(200), primary_key=True)
    email_id = _app.db.Column(_app.db.String(200), nullable=True, unique=True, index=True)
    event_type = _app.db.Column(_app.db.String(80), nullable=False, index=True)
    received_at = _app.db.Column(
        _app.db.DateTime, nullable=False,
        default=lambda: _app.datetime.now(_app.timezone.utc).replace(tzinfo=None))


class InboundInboxMessage(_app.db.Model):
    """Email sent directly to Yazory's public information address."""
    __tablename__ = 'inbound_inbox_message'
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    provider_message_id = _app.db.Column(
        _app.db.String(200), nullable=False, unique=True, index=True)
    sender_name = _app.db.Column(_app.db.String(200), nullable=False, default='')
    sender_email = _app.db.Column(_app.db.String(254), nullable=False, index=True)
    recipient = _app.db.Column(_app.db.String(254), nullable=False)
    subject = _app.db.Column(_app.db.String(300), nullable=False, default='')
    body = _app.db.Column(_app.db.Text, nullable=False, default='')
    html_body = _app.db.Column(_app.db.Text, nullable=False, default='')
    status = _app.db.Column(
        _app.db.String(20), nullable=False, default='unread', index=True)
    created_at = _app.db.Column(
        _app.db.DateTime, nullable=False,
        default=lambda: _app.datetime.now(_app.timezone.utc).replace(tzinfo=None),
        index=True)
    handled_at = _app.db.Column(_app.db.DateTime, nullable=True)
    handled_by = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('staff_user.id'), nullable=True)
    handler = _app.db.relationship('StaffUser')


class GeneralSmsMessage(_app.db.Model):
    """An SMS conversation outside the supporter communication workflow."""
    __tablename__ = 'general_sms_message'
    id = _app.db.Column(_app.db.Integer, primary_key=True)
    provider_message_id = _app.db.Column(
        _app.db.String(100), nullable=True, unique=True, index=True)
    phone = _app.db.Column(_app.db.String(80), nullable=False, index=True)
    direction = _app.db.Column(_app.db.String(20), nullable=False, index=True)
    body = _app.db.Column(_app.db.Text, nullable=False, default='')
    status = _app.db.Column(
        _app.db.String(20), nullable=False, default='unread', index=True)
    delivery_error = _app.db.Column(_app.db.Text, nullable=False, default='')
    staff_user_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('staff_user.id'), nullable=True, index=True)
    family_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('family.id'), nullable=True, index=True)
    created_at = _app.db.Column(
        _app.db.DateTime, nullable=False,
        default=lambda: _app.datetime.now(_app.timezone.utc).replace(tzinfo=None),
        index=True)
    handled_at = _app.db.Column(_app.db.DateTime, nullable=True)
    staff_user = _app.db.relationship('StaffUser')
    family = _app.db.relationship('Family')


class SmsPhoneLink(_app.db.Model):
    """A staff-confirmed link between an SMS number and a directory person."""
    __tablename__ = 'sms_phone_link'
    normalized_phone = _app.db.Column(_app.db.String(20), primary_key=True)
    profile_id = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('supporter_profile.id'),
        nullable=False, index=True)
    created_at = _app.db.Column(
        _app.db.DateTime, nullable=False,
        default=lambda: _app.datetime.now(_app.timezone.utc).replace(tzinfo=None))
    created_by = _app.db.Column(
        _app.db.Integer, _app.db.ForeignKey('staff_user.id'), nullable=True)
    profile = _app.db.relationship('SupporterProfile')


from app_original import *  # noqa: F401,F403,E402
from native_payments import register_native_payments  # noqa: E402
from supporter_portal import register_supporter_portal  # noqa: E402
from applicant_portal import ApplicantMessage, register_applicant_portal  # noqa: E402
from ai_email import (draft_initial_email, fallback_initial_email,
                      initial_email_subject)  # noqa: E402
from inbound_email import (html_to_text, retrieve_received_email,
                           verify_webhook)  # noqa: E402


def _normalize_rabbi_name(name):
    return ' '.join((name or '').split()).casefold()


def _canonical_rabbi(name, phone=''):
    name = ' '.join((name or '').split())[:160]
    if not name:
        return None
    normalized = _normalize_rabbi_name(name)
    person = _app.db.session.scalar(select(RabbiPerson).where(
        RabbiPerson.normalized_name == normalized))
    if person is None:
        person = RabbiPerson(name=name, normalized_name=normalized, phone=(phone or '')[:80])
        _app.db.session.add(person)
        _app.db.session.flush()
    elif phone and person.phone != phone[:80]:
        person.phone = phone[:80]
    return person


def _canonical_helper(name, phones=()):
    name = ' '.join((name or '').split())[:160]
    if not name:
        return None
    normalized = _normalize_rabbi_name(name)
    person = _app.db.session.scalar(select(HelperPerson).where(
        HelperPerson.normalized_name == normalized))
    if person is None:
        person = HelperPerson(name=name, normalized_name=normalized)
        _app.db.session.add(person)
        _app.db.session.flush()
    else:
        person.name = name
    saved = {row.phone.casefold() for row in _app.db.session.scalars(
        select(HelperPhone).where(HelperPhone.helper_person_id == person.id)).all()}
    for phone in _clean_phones(phones):
        if phone.casefold() not in saved:
            _app.db.session.add(HelperPhone(helper_person_id=person.id, phone=phone))
            saved.add(phone.casefold())
    return person


def _attach_helper(institution, person, role):
    association = _app.db.session.scalar(select(ShulHelperAssociation).where(
        ShulHelperAssociation.institution_id == institution.id,
        ShulHelperAssociation.helper_person_id == person.id,
        ShulHelperAssociation.role == role))
    if association is None:
        association = ShulHelperAssociation(
            institution_id=institution.id, helper_person_id=person.id, role=role)
        _app.db.session.add(association)
        _app.db.session.flush()
    return association


def _sync_helper_associations(institution, role, people):
    current = _app.db.session.scalars(select(ShulHelperAssociation).where(
        ShulHelperAssociation.institution_id == institution.id,
        ShulHelperAssociation.role == role)).all()
    desired_ids = {person.id for person in people}
    for association in current:
        if association.helper_person_id not in desired_ids:
            _app.db.session.delete(association)
    for person in people:
        _attach_helper(institution, person, role)


def _helper_phones(person):
    return [row.phone for row in _app.db.session.scalars(select(HelperPhone).where(
        HelperPhone.helper_person_id == person.id).order_by(HelperPhone.id)).all()]


def _helper_payloads(institution, role):
    """Read helper directory data only from the canonical person links."""
    if institution is None:
        return []
    associations = _app.db.session.scalars(select(ShulHelperAssociation).where(
        ShulHelperAssociation.institution_id == institution.id,
        ShulHelperAssociation.role == role
    ).order_by(ShulHelperAssociation.id)).all()
    result = []
    for association in associations:
        phones = _helper_phones(association.helper_person)
        result.append({
            'name': association.helper_person.name,
            'phone': phones[0] if phones else '',
            'phones': phones,
        })
    return result


def _primary_association(institution_id):
    return _app.db.session.scalar(select(ShulRabbiAssociation).where(
        ShulRabbiAssociation.institution_id == institution_id,
        ShulRabbiAssociation.is_primary.is_(True)))


def _sync_legacy_primary(institution_id):
    primary = _primary_association(institution_id)
    legacy = _app.db.session.get(ShulRabbi, institution_id)
    if primary is None:
        if legacy is not None:
            _app.db.session.delete(legacy)
        return
    if legacy is None:
        legacy = ShulRabbi(institution_id=institution_id)
        _app.db.session.add(legacy)
    legacy.rabbi_name = primary.rabbi_person.name
    legacy.rabbi_phone = primary.rabbi_person.phone or ''


def _attach_rabbi(institution, person, make_primary=False):
    association = _app.db.session.scalar(select(ShulRabbiAssociation).where(
        ShulRabbiAssociation.institution_id == institution.id,
        ShulRabbiAssociation.rabbi_person_id == person.id))
    if association is None:
        association = ShulRabbiAssociation(
            institution_id=institution.id, rabbi_person_id=person.id, is_primary=False)
        _app.db.session.add(association)
        _app.db.session.flush()
    if make_primary or _primary_association(institution.id) is None:
        for row in _app.db.session.scalars(select(ShulRabbiAssociation).where(
                ShulRabbiAssociation.institution_id == institution.id)).all():
            row.is_primary = False
        _app.db.session.flush()
        association.is_primary = True
        _app.db.session.flush()
    _sync_legacy_primary(institution.id)
    return association


def _migrate_canonical_rabbis():
    marker_key = 'canonical_rabbis_v1'
    if _app.db.session.get(_app.OrganizationSetting, marker_key) is not None:
        return
    for legacy in _app.db.session.scalars(select(ShulRabbi).order_by(
            ShulRabbi.institution_id)).all():
        person = _canonical_rabbi(legacy.rabbi_name, legacy.rabbi_phone)
        if person is not None:
            institution = _app.db.session.get(_app.Institution, legacy.institution_id)
            if institution is not None:
                _attach_rabbi(institution, person, make_primary=True)
    for family in _app.db.session.scalars(select(_app.Family).order_by(_app.Family.id)).all():
        preference = _app.db.session.get(FamilyRabbiPreference, family.id)
        if preference is None:
            person = _canonical_rabbi(family.rabbi, family.rabbi_phone)
            _app.db.session.add(FamilyRabbiPreference(
                family_id=family.id,
                rabbi_person_id=person.id if person else None,
                overridden=bool(family.rabbi or family.rabbi_phone)))
    _app.db.session.add(_app.OrganizationSetting(key=marker_key, value={'completed': True}))
    _app.db.session.commit()


def _migrate_canonical_helpers():
    marker_key = 'canonical_helpers_v1'
    marker = _app.db.session.get(_app.OrganizationSetting, marker_key)
    for row in _app.db.session.scalars(select(ShulGabbaiDirectory).order_by(
            ShulGabbaiDirectory.id)).all():
        person = _canonical_helper(row.name, _gabbai_phones(row))
        if person is not None:
            _attach_helper(row.institution, person, 'shul_gabbai')
    for row in _app.db.session.scalars(select(ShulRabbiAssistant).order_by(
            ShulRabbiAssistant.id)).all():
        phones = [phone.phone for phone in _app.db.session.scalars(
            select(ShulRabbiAssistantPhone).where(
                ShulRabbiAssistantPhone.assistant_id == row.id
            ).order_by(ShulRabbiAssistantPhone.id)).all()]
        person = _canonical_helper(row.name, phones)
        if person is not None:
            _attach_helper(row.institution, person, 'rabbi_assistant')
    if marker is None:
        _app.db.session.add(_app.OrganizationSetting(
            key=marker_key, value={'completed': True}))
    _app.db.session.commit()


def _family_profile_shuls(family):
    if family is None:
        return []
    selected_names = {
        (name or '').strip().casefold()
        for name in (family.weekday_shul, family.shabbos_shul) if (name or '').strip()
    }
    rows = _app.db.session.scalars(select(_app.Institution).join(
        _app.PersonAffiliation,
        _app.PersonAffiliation.institution_id == _app.Institution.id
    ).where(
        _app.PersonAffiliation.person_type == 'family',
        _app.PersonAffiliation.person_id == family.id,
        _app.Institution.kind == 'Shul',
    ).order_by(_app.Institution.id)).all()
    return [row for row in rows if row.name.strip().casefold() in selected_names]


def _migrate_family_gabbaim_to_shared_shuls():
    """Move unambiguous legacy family gabbaim onto their one shared shul."""
    marker_key = 'shared_shul_gabbaim_v2'
    if _app.db.session.get(_app.OrganizationSetting, marker_key) is not None:
        return
    for family in _app.db.session.scalars(select(_app.Family).order_by(_app.Family.id)).all():
        shuls = _family_profile_shuls(family)
        if len(shuls) != 1:
            continue
        institution = shuls[0]
        for legacy in list(family.gabbais):
            phones = _clean_phones([legacy.phone])
            person = _canonical_helper(legacy.name, phones)
            if person is not None:
                _attach_helper(institution, person, 'shul_gabbai')
            _app.db.session.delete(legacy)
    # These copied rows represented a second connection path. Family profiles
    # now derive gabbaim only through Family -> shared Institution.
    _app.db.session.execute(_app.db.delete(FamilyGabbaiConnection))
    _app.db.session.add(_app.OrganizationSetting(key=marker_key, value={'completed': True}))
    _app.db.session.commit()


def _submitted_institution_name(key):
    value = _app.request.form.get(key, '').strip()
    if value == '__new__':
        value = _app.request.form.get(key + '_new', '').strip()
    return value[:160]


def _family_id_from_response(response):
    location = getattr(response, 'location', None) or response.headers.get('Location', '')
    match = re.search(r'/families/(\d+)(?:$|[/?#])', location)
    return int(match.group(1)) if match else None


def _find_shul(shul_name):
    if not shul_name:
        return None
    return _app.db.session.scalar(select(_app.Institution).where(
        _app.Institution.kind == 'Shul',
        _app.Institution.name == shul_name).order_by(_app.Institution.id))


def _clean_phones(values):
    result = []
    seen = set()
    for raw in values:
        phone = (raw or '').strip()[:80]
        if not phone or phone.casefold() in seen:
            continue
        seen.add(phone.casefold())
        result.append(phone)
    return result


def _family_phones(family):
    if family is None:
        return []
    rows = _app.db.session.scalars(select(FamilyPhone).where(
        FamilyPhone.family_id == family.id).order_by(FamilyPhone.id)).all()
    phones = [row.phone for row in rows]
    if family.phone and family.phone not in phones:
        phones.insert(0, family.phone)
    return phones


def _replace_family_phones(family_id, phones):
    cleaned = _clean_phones(phones)
    family = _app.db.session.get(_app.Family, family_id)
    if family is None:
        return
    family.phone = cleaned[0] if cleaned else ''
    current = _app.db.session.scalars(select(FamilyPhone).where(
        FamilyPhone.family_id == family_id)).all()
    for row in current:
        _app.db.session.delete(row)
    _app.db.session.flush()
    for phone in cleaned:
        _app.db.session.add(FamilyPhone(family_id=family_id, phone=phone))


def _json_phone_lists(field):
    result = []
    for raw in _app.request.form.getlist(field):
        try:
            values = json.loads(raw or '[]')
        except (TypeError, ValueError):
            values = []
        result.append(_clean_phones(values if isinstance(values, list) else []))
    return result


def _rabbi_phones(institution):
    rows = _app.db.session.scalars(select(ShulRabbiPhone).where(
        ShulRabbiPhone.institution_id == institution.id).order_by(ShulRabbiPhone.id)).all()
    if rows:
        return [row.phone for row in rows]
    legacy = _app.db.session.get(ShulRabbi, institution.id)
    return [legacy.rabbi_phone] if legacy and legacy.rabbi_phone else []


def _gabbai_phones(gabbai):
    rows = _app.db.session.scalars(select(ShulGabbaiPhone).where(
        ShulGabbaiPhone.gabbai_id == gabbai.id).order_by(ShulGabbaiPhone.id)).all()
    if rows:
        return [row.phone for row in rows]
    return [gabbai.phone] if gabbai.phone else []


def _replace_rabbi_phones(institution, phones):
    current = _app.db.session.scalars(select(ShulRabbiPhone).where(
        ShulRabbiPhone.institution_id == institution.id)).all()
    for row in current:
        _app.db.session.delete(row)
    _app.db.session.flush()
    for phone in _clean_phones(phones):
        _app.db.session.add(ShulRabbiPhone(institution_id=institution.id, phone=phone))


def _save_primary_rabbi_phone(person, phone):
    """Keep a phone entered on a family form in the shared rabbi directory."""
    phone = (phone or '').strip()[:80]
    if person is None or not phone:
        return
    person.phone = phone
    associations = _app.db.session.scalars(select(ShulRabbiAssociation).where(
        ShulRabbiAssociation.rabbi_person_id == person.id,
        ShulRabbiAssociation.is_primary.is_(True))).all()
    for association in associations:
        institution = association.institution
        phones = _rabbi_phones(institution)
        _replace_rabbi_phones(
            institution, [phone] + [saved for saved in phones if saved != phone])
        _sync_legacy_primary(institution.id)


def _replace_gabbai_phones(gabbai, phones):
    current = _app.db.session.scalars(select(ShulGabbaiPhone).where(
        ShulGabbaiPhone.gabbai_id == gabbai.id)).all()
    for row in current:
        _app.db.session.delete(row)
    _app.db.session.flush()
    for phone in _clean_phones(phones):
        _app.db.session.add(ShulGabbaiPhone(gabbai_id=gabbai.id, phone=phone))


def _upsert_family_rabbi(family_id, role, institution, rabbi_name, rabbi_phone):
    row = _app.db.session.scalar(select(FamilyRabbiConnection).where(
        FamilyRabbiConnection.family_id == family_id,
        FamilyRabbiConnection.role == role))
    if not rabbi_name:
        if row is not None:
            _app.db.session.delete(row)
        return
    if row is None:
        row = FamilyRabbiConnection(family_id=family_id, role=role)
        _app.db.session.add(row)
    row.institution_id = institution.id if institution else None
    row.rabbi_name = rabbi_name
    row.rabbi_phone = rabbi_phone


def _replace_rabbi_assistants(institution, names, phone_lists):
    helper_people = []
    for index, raw_name in enumerate(names):
        name = (raw_name or '').strip()[:160]
        if not name:
            continue
        phones = phone_lists[index] if index < len(phone_lists) else []
        helper = _canonical_helper(name, phones)
        if helper is not None:
            helper_people.append(helper)
    _sync_helper_associations(institution, 'rabbi_assistant', helper_people)


def _save_shul_rabbi_connections(family_id):
    family = _app.db.session.get(_app.Family, family_id)
    if family is None:
        return

    submitted = []
    for role, key in (('weekday_shul', 'weekday_shul'), ('shabbos_shul', 'shabbos_shul')):
        shul_name = _submitted_institution_name(key)
        if not shul_name:
            _upsert_family_rabbi(family_id, role, None, '', '')
            continue
        institution = _find_shul(shul_name)
        can_manage_directory = current_app.config['DEMO']
        if not can_manage_directory:
            user = _app.db.session.get(_app.StaffUser, session.get('user_id'))
            can_manage_directory = bool(user and user.role == 'organization_admin')
        rabbi_name = (_app.request.form.get(key + '_rabbi', '').strip()[:160]
                      if can_manage_directory else '')
        phone_field = key + '_rabbi_phone'
        has_phone_fields = can_manage_directory and phone_field in _app.request.form
        phones = (_clean_phones(_app.request.form.getlist(phone_field))
                  if can_manage_directory else [])
        assistant_field = key + '_rabbi_assistant_name'
        has_assistant_fields = can_manage_directory and assistant_field in _app.request.form
        assistant_names = (_app.request.form.getlist(assistant_field)
                           if can_manage_directory else [])
        assistant_phones = (_json_phone_lists(key + '_rabbi_assistant_phones')
                            if can_manage_directory else [])
        submitted.append((
            role, institution, rabbi_name, phones, has_phone_fields,
            assistant_names, assistant_phones, has_assistant_fields))

    assignments = {}
    inherited_rabbis = []
    for (role, institution, rabbi_name, submitted_phones, has_phone_fields,
         assistant_names, assistant_phones, has_assistant_fields) in submitted:
        if institution is None:
            first_phone = submitted_phones[0] if submitted_phones else ''
            _upsert_family_rabbi(family_id, role, None, rabbi_name, first_phone)
            if rabbi_name:
                inherited_rabbis.append((rabbi_name, first_phone, None))
            continue

        existing = _app.db.session.get(ShulRabbi, institution.id)
        chosen = assignments.get(institution.id)
        if chosen is None:
            if rabbi_name:
                phones = submitted_phones if has_phone_fields else _rabbi_phones(institution)
                chosen = (rabbi_name, phones)
            elif existing is not None:
                chosen = (existing.rabbi_name, _rabbi_phones(institution))
            else:
                chosen = ('', [])
            assignments[institution.id] = chosen

        rabbi_name, phones = chosen
        first_phone = phones[0] if phones else ''
        if rabbi_name:
            person = _canonical_rabbi(rabbi_name, first_phone)
            _attach_rabbi(institution, person, make_primary=(
                existing is None or existing.rabbi_name != rabbi_name or
                existing.rabbi_phone != first_phone))
            existing = _app.db.session.get(ShulRabbi, institution.id)
            rabbi_name = existing.rabbi_name
            first_phone = existing.rabbi_phone
            if has_phone_fields:
                _replace_rabbi_phones(institution, submitted_phones)
            elif not _rabbi_phones(institution) and first_phone:
                _replace_rabbi_phones(institution, [first_phone])
            if has_assistant_fields:
                _replace_rabbi_assistants(
                    institution, assistant_names, assistant_phones)
            inherited_rabbis.append((rabbi_name, first_phone, institution))
        _upsert_family_rabbi(
            family_id, role, institution, rabbi_name, first_phone)

    if not (family.rabbi or '').strip() and inherited_rabbis:
        unique = {}
        for rabbi_name, rabbi_phone, institution in inherited_rabbis:
            unique.setdefault(rabbi_name.casefold(), (rabbi_name, rabbi_phone, institution))
        if len(unique) == 1:
            rabbi_name, rabbi_phone, institution = next(iter(unique.values()))
            family.rabbi = rabbi_name
            family.rabbi_phone = rabbi_phone
            _upsert_family_rabbi(
                family_id, 'affiliated', institution, rabbi_name, rabbi_phone)
        else:
            _upsert_family_rabbi(
                family_id, 'affiliated', None, family.rabbi or '', family.rabbi_phone or '')
    else:
        _upsert_family_rabbi(
            family_id, 'affiliated', None, family.rabbi or '', family.rabbi_phone or '')


def _save_family_rabbi_preference(family_id):
    family = _app.db.session.get(_app.Family, family_id)
    preference = _app.db.session.get(FamilyRabbiPreference, family_id)
    if preference is None:
        preference = FamilyRabbiPreference(family_id=family_id)
        _app.db.session.add(preference)
    mode = _app.request.form.get('rabbi_mode', '').strip()
    selected_id = _app.request.form.get('rabbi_person_id', type=int)
    if mode == 'automatic':
        preference.overridden = False
        person = None
        selected_institution = None
        for key in ('weekday_shul', 'shabbos_shul'):
            institution = _find_shul(_submitted_institution_name(key))
            primary = _primary_association(institution.id) if institution else None
            if primary:
                person = primary.rabbi_person
                selected_institution = institution
                break
        submitted_name = _app.request.form.get('rabbi', '').strip()
        submitted_phone = _app.request.form.get('rabbi_phone', '').strip()
        if (person is not None and submitted_phone and
                (not submitted_name or
                 _normalize_rabbi_name(submitted_name) == person.normalized_name)):
            _save_primary_rabbi_phone(person, submitted_phone)
        preference.rabbi_person_id = person.id if person else None
        family.rabbi = person.name if person else ''
        family.rabbi_phone = person.phone if person else ''
    elif mode == 'canonical':
        person = _app.db.session.get(RabbiPerson, selected_id) if selected_id else None
        if person is None:
            _app.abort(400, 'Choose a valid rabbi.')
        _save_primary_rabbi_phone(
            person, _app.request.form.get('rabbi_phone', '').strip())
        preference.overridden = True
        preference.rabbi_person_id = person.id
        family.rabbi, family.rabbi_phone = person.name, person.phone
    else:
        name = _app.request.form.get('rabbi', '').strip()[:160]
        phone = _app.request.form.get('rabbi_phone', '').strip()[:80]
        # Legacy clients with a nonblank affiliated rabbi are manual overrides;
        # blank legacy submissions retain the historical automatic behavior.
        manual = mode == 'manual' or bool(name or phone)
        if manual:
            person = _canonical_rabbi(name, phone)
            preference.overridden = True
            preference.rabbi_person_id = person.id if person else None
            family.rabbi, family.rabbi_phone = name, phone
        else:
            preference.overridden = False
            person = None
            selected_institution = None
            for key in ('weekday_shul', 'shabbos_shul'):
                institution = _find_shul(_submitted_institution_name(key))
                primary = _primary_association(institution.id) if institution else None
                if primary:
                    person = primary.rabbi_person
                    selected_institution = institution
                    break
            preference.rabbi_person_id = person.id if person else None
            family.rabbi = person.name if person else ''
            family.rabbi_phone = person.phone if person else ''
    affiliated_institution = (selected_institution
                              if not preference.overridden else None)
    _upsert_family_rabbi(
        family_id, 'affiliated', affiliated_institution,
        family.rabbi or '', family.rabbi_phone or '')


def _sync_automatic_family_rabbis():
    """Refresh role snapshots and non-overridden family projections."""
    for family in _app.db.session.scalars(select(_app.Family)).all():
        selected_person = None
        for role, name in (('weekday_shul', family.weekday_shul),
                           ('shabbos_shul', family.shabbos_shul)):
            institution = _find_shul(name)
            primary = _primary_association(institution.id) if institution else None
            person = primary.rabbi_person if primary else None
            _upsert_family_rabbi(
                family.id, role, institution,
                person.name if person else '', person.phone if person else '')
            if selected_person is None and person is not None:
                selected_person = person
        preference = _app.db.session.get(FamilyRabbiPreference, family.id)
        if preference is not None and not preference.overridden:
            preference.rabbi_person_id = selected_person.id if selected_person else None
            family.rabbi = selected_person.name if selected_person else ''
            family.rabbi_phone = selected_person.phone if selected_person else ''
            _upsert_family_rabbi(
                family.id, 'affiliated', None, family.rabbi, family.rabbi_phone)


def _submitted_gabbais(key):
    names = _app.request.form.getlist(key + '_gabbai_name')
    phone_lists = _json_phone_lists(key + '_gabbai_phones')
    legacy_phones = _app.request.form.getlist(key + '_gabbai_phone')
    rows = []
    seen = set()
    for index, raw_name in enumerate(names):
        name = (raw_name or '').strip()[:160]
        if not name:
            continue
        phones = phone_lists[index] if index < len(phone_lists) else []
        if not phones and index < len(legacy_phones):
            phones = _clean_phones([legacy_phones[index]])
        identity = name.casefold()
        if identity in seen:
            continue
        seen.add(identity)
        rows.append((name, phones))
    return rows


def _save_shul_gabbai_connections(family_id):
    submitted_by_institution = {}
    for role, key in (('weekday_shul', 'weekday_shul'), ('shabbos_shul', 'shabbos_shul')):
        shul_name = _submitted_institution_name(key)
        institution = _find_shul(shul_name)
        submitted = _submitted_gabbais(key)
        if (institution is not None and
                institution.id not in submitted_by_institution and
                key + '_gabbai_name' in _app.request.form):
            submitted_by_institution[institution.id] = submitted

    for institution_id, rows in submitted_by_institution.items():
        institution = _app.db.session.get(_app.Institution, institution_id)
        people = []
        for name, phones in rows:
            person = _canonical_helper(name, phones)
            if person is not None:
                people.append(person)
        _sync_helper_associations(institution, 'shul_gabbai', people)


def _save_shul_connections(family_id):
    _save_shul_rabbi_connections(family_id)
    _sync_automatic_family_rabbis()
    _save_family_rabbi_preference(family_id)
    _save_shul_gabbai_connections(family_id)
    _app.db.session.commit()


def _assistant_payload(institution):
    return _helper_payloads(institution, 'rabbi_assistant')


def create_app(test_config=None):
    app = _app.create_app(test_config)

    # Core supporter forms live in app_original.py, while the imported people
    # directory is supplied by this compatibility layer. Publish the model to
    # the core routes without introducing an import cycle.
    app.extensions['supporter_profile_model'] = SupporterProfile

    def imported_supporter_profiles():
        return _app.db.session.scalars(select(SupporterProfile).order_by(
            SupporterProfile.name, SupporterProfile.id)).all()

    app.jinja_env.globals['imported_supporter_profiles'] = imported_supporter_profiles

    @app.get('/supporter-directory/options')
    def supporter_directory_options():
        supporter_directory_families()
        profiles = imported_supporter_profiles()
        return jsonify({
            'profiles': [{
                'id': profile.id, 'name': profile.name,
                'phone': profile.phone, 'email': profile.email,
            } for profile in profiles],
            'labels': {
                'choose': _app.translate('Choose from imported people'),
                'hint': _app.translate('Search by name, phone, or email'),
                'search': _app.translate('Search people'),
                'new': _app.translate('Enter a new person below'),
            },
        })

    def extended_directory_people():
        """Expose every reusable person table through the shared person picker."""
        people = []
        for person in _app.db.session.scalars(select(SupporterProfile).order_by(
                SupporterProfile.name)).all():
            people.append(('supporter_profile', person.id, person.name,
                           'Supporter', person.email or person.phone))
        for person in _app.db.session.scalars(select(RabbiPerson).order_by(
                RabbiPerson.name)).all():
            people.append(('rabbi', person.id, person.name, 'Rabbi', person.phone))
        helper_phones = {row.helper_person_id: row.phone for row in
                         _app.db.session.scalars(select(HelperPhone).order_by(
                             HelperPhone.id)).all()}
        for person in _app.db.session.scalars(select(HelperPerson).order_by(
                HelperPerson.name)).all():
            people.append(('helper', person.id, person.name, 'Community helper',
                           helper_phones.get(person.id, '')))
        return people

    app.extensions.setdefault('person_directory_providers', []).append(
        extended_directory_people)

    def resolve_extended_directory_person(person_type, person_id):
        model = {
            'supporter_profile': SupporterProfile,
            'rabbi': RabbiPerson,
            'helper': HelperPerson,
        }.get(person_type)
        person = _app.db.session.get(model, person_id) if model else None
        if person is None:
            return None
        phone = getattr(person, 'phone', '')
        if person_type == 'helper':
            phone = _app.db.session.scalar(select(HelperPhone.phone).where(
                HelperPhone.helper_person_id == person.id).order_by(HelperPhone.id)) or ''
        return {'name': person.name, 'phone': phone,
                'email': getattr(person, 'email', '')}

    app.extensions.setdefault('person_directory_resolvers', []).append(
        resolve_extended_directory_person)

    def create_neutral_directory_person(name, phone='', email=''):
        """Create the neutral identity once; roles can be attached later."""
        normalized = normalized_profile_phone(phone)
        if normalized:
            existing = _app.db.session.scalar(select(SupporterProfile).where(
                SupporterProfile.normalized_phone == normalized))
            if existing:
                if not existing.email and email:
                    existing.email = email[:254]
                return existing
        else:
            normalized = 'person:' + _app.secrets.token_hex(12)
        person = SupporterProfile(name=name[:160], phone=phone[:80],
                                  normalized_phone=normalized,
                                  email=email[:254])
        _app.db.session.add(person)
        return person

    app.extensions.setdefault('person_directory_creators', []).append(
        create_neutral_directory_person)

    family_relative_fields = (
        ('name', 'applicant', 'Applicant'),
        ('spouse', 'spouse', 'Spouse'),
        ('father', 'father', 'Father'),
        ('inlaws', 'inlaws', 'Father-in-law'),
        ('inlaws_maiden_name', 'maiden', "Father’s father-in-law"),
        ('inlaws_family', 'inlawfam', "Father-in-law’s father-in-law"),
    )

    def sync_family_relatives_to_people(family):
        """Publish names entered on a case in the independent people list.

        The deterministic directory key is important for relatives who do not
        yet have a phone number: release-time backfills and later profile saves
        update one record instead of creating a duplicate every time.
        """
        for field_name, key_suffix, relationship in family_relative_fields:
            name = (getattr(family, field_name, '') or '').strip()[:160]
            if not name:
                continue
            directory_key = f'family:{family.id}:{key_suffix}'
            profile = _app.db.session.scalar(select(SupporterProfile).where(
                SupporterProfile.normalized_phone == directory_key))
            if profile is None:
                profile = SupporterProfile(
                    name=name,
                    phone=(family.phone or '')[:80] if field_name == 'name' else '',
                    normalized_phone=directory_key,
                    email=(family.email or '')[:254] if field_name == 'name' else '')
                _app.db.session.add(profile)
            else:
                profile.name = name
                if field_name == 'name':
                    profile.phone = (family.phone or '')[:80]
                    profile.email = (family.email or '')[:254]
            _app.db.session.flush()
            person = canonical_person_for_profile(profile)
            person.name = name
            if field_name == 'name':
                person.phone = (family.phone or '')[:80]
                person.cell_phone = (family.phone or '')[:80]
                person.email = (family.email or '')[:254]
                person.home_address = (family.address or '')[:240]
                person.city = (family.city or '')[:120]
                person.state = (family.state or '')[:80]
                person.zip_code = (family.zip_code or '')[:20]
            source_note = f'{relationship} of {family.name} (YZ-{family.id:04d})'
            generated_prefixes = tuple(
                label for _, _, label in family_relative_fields)
            if not person.notes or person.notes.startswith(generated_prefixes):
                person.notes = source_note

    app.extensions.setdefault('family_profile_person_sync', []).append(
        sync_family_relatives_to_people)

    def sync_askan_to_people(askan):
        """Keep every askan in the same canonical people directory."""
        _app.db.session.flush()
        directory_key = f'askan:{askan.id}'
        profile = _app.db.session.scalar(select(SupporterProfile).where(
            SupporterProfile.normalized_phone == directory_key))
        if profile is None:
            normalized = normalized_profile_phone(askan.phone)
            if normalized:
                profile = _app.db.session.scalar(select(SupporterProfile).where(
                    SupporterProfile.normalized_phone == normalized))
            if profile is None and askan.email:
                profile = _app.db.session.scalar(select(SupporterProfile).where(
                    _app.func.lower(SupporterProfile.email) == askan.email.lower()))
        if profile is None:
            profile = SupporterProfile(
                name=askan.name, phone=askan.phone or '',
                normalized_phone=directory_key, email=askan.email or '')
            _app.db.session.add(profile)
        else:
            profile.name = askan.name
            profile.phone = askan.phone or ''
            profile.email = askan.email or ''
        _app.db.session.flush()
        person = canonical_person_for_profile(profile)
        person.name = askan.name
        person.phone = askan.phone or ''
        person.cell_phone = askan.phone or ''
        person.email = askan.email or ''
        if not person.notes or person.notes.startswith('Askan in Yazory directory'):
            person.notes = 'Askan in Yazory directory'

    app.extensions.setdefault('askan_profile_person_sync', []).append(
        sync_askan_to_people)
    twilio_config = {
        'TWILIO_ACCOUNT_SID': os.getenv('TWILIO_ACCOUNT_SID', ''),
        'TWILIO_AUTH_TOKEN': os.getenv('TWILIO_AUTH_TOKEN', ''),
        'TWILIO_SMS_FROM': os.getenv('TWILIO_SMS_FROM', ''),
        'TWILIO_MESSAGING_SERVICE_SID': os.getenv('TWILIO_MESSAGING_SERVICE_SID', ''),
        'TWILIO_WHATSAPP_FROM': os.getenv('TWILIO_WHATSAPP_FROM', ''),
    }
    if test_config:
        twilio_config.update({key: test_config[key] for key in twilio_config if key in test_config})
    app.config.update(twilio_config)

    personal_fields = (
        'name', 'phone', 'email', 'home_phone', 'cell_phone',
        'home_address', 'city', 'state', 'zip_code', 'workplace',
        'work_phone', 'notes')

    def sync_person_snapshots(person):
        """Maintain old Contact readers while SupporterPerson is authoritative."""
        contacts = _app.db.session.scalars(select(_app.Contact).where(
            _app.Contact.person_id == person.id)).all()
        for row in contacts:
            for field_name in personal_fields:
                setattr(row, field_name, getattr(person, field_name) or '')
            row.supporter_key = person.identity_key

        profile = _app.db.session.scalar(select(SupporterProfile).where(
            SupporterProfile.person_id == person.id))
        if profile is not None:
            profile.name = person.name
            profile.phone = person.phone
            profile.email = person.email
            normalized = normalized_profile_phone(person.phone)
            if normalized:
                profile.normalized_phone = normalized

    def canonical_person_for_profile(profile):
        """Attach an import row to the same authoritative person used by cases."""
        person = (_app.db.session.get(SupporterPerson, profile.person_id)
                  if profile.person_id else None)
        if person is None:
            identity_key = 'phone:' + profile.normalized_phone
            person = _app.db.session.scalar(select(SupporterPerson).where(
                SupporterPerson.identity_key == identity_key))
            if person is None:
                person = SupporterPerson(
                    identity_key=identity_key, name=profile.name,
                    phone=profile.phone, cell_phone=profile.phone,
                    email=profile.email)
                _app.db.session.add(person)
                _app.db.session.flush()
            else:
                if not person.email:
                    person.email = profile.email
                if not person.cell_phone:
                    person.cell_phone = profile.phone
            profile.person_id = person.id
        return person

    def canonicalize_unlinked_profiles():
        """Make imported profiles selectable using a fixed number of queries."""
        profiles = _app.db.session.scalars(select(SupporterProfile).where(
            SupporterProfile.person_id.is_(None))).all()
        if not profiles:
            return
        identity_keys = ['phone:' + row.normalized_phone for row in profiles]
        people_by_key = {
            row.identity_key: row for row in _app.db.session.scalars(select(
                SupporterPerson).where(
                    SupporterPerson.identity_key.in_(identity_keys))).all()
        }
        profile_people = []
        for profile, identity_key in zip(profiles, identity_keys):
            person = people_by_key.get(identity_key)
            if person is None:
                person = SupporterPerson(
                    identity_key=identity_key, name=profile.name,
                    phone=profile.phone, cell_phone=profile.phone,
                    email=profile.email)
                _app.db.session.add(person)
                people_by_key[identity_key] = person
            else:
                if not person.email:
                    person.email = profile.email
                if not person.cell_phone:
                    person.cell_phone = profile.phone
            profile_people.append((profile, person))
        _app.db.session.flush()
        for profile, person in profile_people:
            profile.person_id = person.id

    def attach_supporter_person(contact, source=None):
        """Attach a case connection to exactly one canonical person."""
        if source is not None and source.person_id is None:
            attach_supporter_person(source)
        if source is not None and source.person_id is not None:
            contact.person_id = source.person_id
            person = _app.db.session.get(SupporterPerson, source.person_id)
        elif contact.person_id is not None:
            person = _app.db.session.get(SupporterPerson, contact.person_id)
        else:
            identity_key = contact.supporter_key or f'legacy:{contact.id}'
            person = _app.db.session.scalar(select(SupporterPerson).where(
                SupporterPerson.identity_key == identity_key))
            # A household phone is not proof that two relatives in the same
            # case are one person. Across cases it is the intended shared
            # identity; within one case preserve distinct people.
            if person is not None and _app.db.session.scalar(select(_app.Contact.id).where(
                    _app.Contact.person_id == person.id,
                    _app.Contact.family_id == contact.family_id,
                    _app.Contact.id != contact.id)):
                identity_key = f'legacy:{contact.id}'
                person = None
            if person is None:
                person = SupporterPerson(
                    identity_key=identity_key,
                    **{field_name: (getattr(contact, field_name, '') or '')
                       for field_name in personal_fields})
                _app.db.session.add(person)
                _app.db.session.flush()
            else:
                # Backfills may encounter an older sparse case copy first.
                for field_name in personal_fields:
                    if not getattr(person, field_name) and getattr(contact, field_name, ''):
                        setattr(person, field_name, getattr(contact, field_name))
            contact.person_id = person.id
            if not identity_key.startswith('legacy:'):
                represented_families = set(_app.db.session.scalars(select(
                    _app.Contact.family_id).where(
                        _app.Contact.person_id == person.id)).all())
                # Repair unsynchronized legacy copies of this identity, but
                # never collapse two people within the same case.
                candidates = _app.db.session.scalars(select(_app.Contact).where(
                    _app.Contact.person_id.is_(None),
                    _app.Contact.supporter_key == identity_key,
                    _app.Contact.id != contact.id).order_by(_app.Contact.id)).all()
                represented_families.add(contact.family_id)
                for candidate in candidates:
                    if candidate.family_id not in represented_families:
                        candidate.person_id = person.id
                        represented_families.add(candidate.family_id)
        sync_person_snapshots(person)
        return person

    def sync_shul_gabbai_supporters(family_ids=None):
        """Connect shared shul gabbaim to each case's Circle of Support.

        Existing case contacts are retained, including their pledge and status.
        The shul link remains the source of the gabbai's role.
        """
        link_model = app.extensions['workflows']['models']['SupporterLink']
        families = _app.db.session.scalars(select(_app.Family).where(
            _app.Family.id.in_(family_ids)) if family_ids is not None else
            select(_app.Family)).all()
        for family in families:
            existing = _app.db.session.scalars(select(_app.Contact).where(
                _app.Contact.family_id == family.id)).all()
            for shul in _family_profile_shuls(family):
                for association in _app.db.session.scalars(select(
                        ShulHelperAssociation).where(
                        ShulHelperAssociation.institution_id == shul.id,
                        ShulHelperAssociation.role == 'shul_gabbai')).all():
                    helper = association.helper_person
                    phones = _helper_phones(helper)
                    phone = phones[0] if phones else ''
                    normalized = normalized_profile_phone(phone)
                    key = 'phone:' + normalized if normalized else f'helper:{helper.id}'
                    shared = _app.db.session.scalar(select(_app.Contact).where(
                        _app.Contact.supporter_key == key).order_by(_app.Contact.id))
                    if shared and shared.name.strip().casefold() != helper.name.strip().casefold():
                        key = f'helper:{helper.id}'
                    contact = next((row for row in existing if
                        row.supporter_key == key or (
                            row.name.strip().casefold() == helper.name.strip().casefold()
                            and (not normalized or normalized in {
                                normalized_profile_phone(row.phone),
                                normalized_profile_phone(row.cell_phone)}))), None)
                    if contact is None:
                        contact = _app.Contact(
                            family_id=family.id, name=helper.name, phone=phone,
                            cell_phone=phone, relationship='Other',
                            supporter_key=key, status='To contact')
                        _app.db.session.add(contact)
                        _app.db.session.flush()
                        source = _app.db.session.scalar(select(_app.Contact).where(
                            _app.Contact.supporter_key == key,
                            _app.Contact.id != contact.id).order_by(_app.Contact.id))
                        attach_supporter_person(contact, source=source)
                        existing.append(contact)
                    if _app.db.session.get(link_model, contact.id) is None:
                        _app.db.session.add(link_model(
                            contact_id=contact.id, side='Community',
                            relationship='Other', introduced_by=shul.name,
                            permission='Not requested', preference='',
                            verified=False))


    def update_supporter_person(contact, values):
        """Update one person once, then refresh every connected case snapshot."""
        person = attach_supporter_person(contact)
        identity_key = values.pop('supporter_key', person.identity_key)
        collision = _app.db.session.scalar(select(SupporterPerson.id).where(
            SupporterPerson.identity_key == identity_key,
            SupporterPerson.id != person.id))
        if collision:
            raise ValueError('That phone number already belongs to another person.')
        person.identity_key = identity_key
        for field_name in personal_fields:
            if field_name in values:
                setattr(person, field_name, values[field_name] or '')
        sync_person_snapshots(person)
        return person

    app.extensions['supporter_identity'] = {
        'attach': attach_supporter_person,
        'update': update_supporter_person,
        'sync': sync_person_snapshots,
        'profile_person': canonical_person_for_profile,
    }

    def ensure_extension_schema():
        """Create and migrate models supplied by this compatibility layer."""
        _app.db.create_all()
        contact_columns = {
            column['name'] for column in _app.inspect(_app.db.engine).get_columns('contact')
        }
        if 'person_id' not in contact_columns:
            _app.db.session.execute(text(
                'ALTER TABLE contact ADD COLUMN person_id INTEGER REFERENCES supporter_person(id)'
            ))
        _app.db.session.execute(text(
            'CREATE INDEX IF NOT EXISTS ix_contact_person_id ON contact (person_id)'
        ))
        profile_columns = {
            column['name'] for column in
            _app.inspect(_app.db.engine).get_columns('supporter_profile')
        }
        if 'person_id' not in profile_columns:
            _app.db.session.execute(text(
                'ALTER TABLE supporter_profile ADD COLUMN person_id INTEGER '
                'REFERENCES supporter_person(id)'))
        _app.db.session.execute(text(
            'CREATE UNIQUE INDEX IF NOT EXISTS uq_supporter_profile_person_id '
            'ON supporter_profile (person_id)'))
        task_columns = {
            column['name'] for column in _app.inspect(_app.db.engine).get_columns('staff_task')
        }
        if 'source_contact_id' not in task_columns:
            _app.db.session.execute(text(
                'ALTER TABLE staff_task ADD COLUMN source_contact_id INTEGER REFERENCES contact(id)'
            ))
        if 'outcome' not in task_columns:
            _app.db.session.execute(text(
                "ALTER TABLE staff_task ADD COLUMN outcome TEXT NOT NULL DEFAULT ''"
            ))
        communication_columns = {
            column['name'] for column in
            _app.inspect(_app.db.engine).get_columns('supporter_communication')
        }
        if 'provider_message_id' not in communication_columns:
            _app.db.session.execute(text(
                "ALTER TABLE supporter_communication ADD COLUMN provider_message_id "
                "VARCHAR(100) NOT NULL DEFAULT ''"))
        if 'delivery_error' not in communication_columns:
            _app.db.session.execute(text(
                "ALTER TABLE supporter_communication ADD COLUMN delivery_error "
                "TEXT NOT NULL DEFAULT ''"))
        inbox_columns = {
            column['name'] for column in
            _app.inspect(_app.db.engine).get_columns('inbound_inbox_message')
        }
        if 'html_body' not in inbox_columns:
            _app.db.session.execute(text(
                "ALTER TABLE inbound_inbox_message ADD COLUMN html_body "
                "TEXT NOT NULL DEFAULT ''"))
        general_sms_columns = {
            column['name'] for column in
            _app.inspect(_app.db.engine).get_columns('general_sms_message')
        }
        if 'family_id' not in general_sms_columns:
            _app.db.session.execute(text(
                'ALTER TABLE general_sms_message ADD COLUMN family_id INTEGER '
                'REFERENCES family(id)'))
        _app.db.session.execute(text(
            'CREATE INDEX IF NOT EXISTS ix_general_sms_message_family_id '
            'ON general_sms_message (family_id)'))
        _app.db.session.execute(text(
            'CREATE UNIQUE INDEX IF NOT EXISTS uq_staff_task_source_contact '
            'ON staff_task (source_contact_id)'
        ))
        _app.db.session.commit()
        # Convert every legacy case copy into a link to one canonical person.
        # This is additive: financial/history rows continue pointing at Contact.
        for contact in _app.db.session.scalars(select(_app.Contact).order_by(
                _app.Contact.id)).all():
            attach_supporter_person(contact)
        _app.db.session.commit()

        # Keep the import staging directory populated for the existing UI.
        existing_profiles = {
            row.normalized_phone: row for row in
            _app.db.session.scalars(select(SupporterProfile)).all()
        }
        for contact in _app.db.session.scalars(select(_app.Contact).where(
                _app.Contact.phone != '')).all():
            normalized = normalized_profile_phone(contact.phone)
            if not normalized:
                continue
            profile = existing_profiles.get(normalized)
            if profile is None:
                profile = SupporterProfile(
                    name=contact.name, phone=contact.phone,
                    normalized_phone=normalized, email=contact.email or '')
                _app.db.session.add(profile)
                existing_profiles[normalized] = profile
            elif not profile.email and contact.email:
                profile.email = contact.email
        # Family relationship fields historically lived only as text on the
        # case. Backfill them during the release migration, so an old case does
        # not have to be opened and saved again.
        for family in _app.db.session.scalars(select(_app.Family)).all():
            sync_family_relatives_to_people(family)
        for askan in _app.db.session.scalars(select(_app.Askan)).all():
            sync_askan_to_people(askan)
        _app.db.session.flush()
        for profile in _app.db.session.scalars(select(SupporterProfile)).all():
            canonical_person_for_profile(profile)
        _app.db.session.commit()
        _migrate_canonical_rabbis()
        _migrate_canonical_helpers()
        # Existing married children become case supporters on the release migration.
        sync_child = app.extensions['sync_married_child_supporters']
        for child in _app.db.session.scalars(select(_app.Child).where(
                _app.Child.married.is_(True))).all():
            sync_child(child)
        _migrate_family_gabbaim_to_shared_shuls()
        sync_shul_gabbai_supporters()
        _app.db.session.commit()

    app.extensions.setdefault('init_db_hooks', []).append(ensure_extension_schema)
    # Keep production worker startup below Replit's health-check deadline.
    # Production receives schema/data maintenance through the explicit
    # `flask --app 'app:create_app()' init-db` release step. Disposable demo
    # and test databases continue to initialize themselves here.
    if app.config['DEMO'] or app.config.get('TESTING'):
        with app.app_context():
            ensure_extension_schema()

    def require_supporter_directory_access():
        user_id = _app.session.get('user_id')
        if app.config['DEMO']:
            return None
        user = _app.db.session.get(_app.StaffUser, user_id) if user_id else None
        if user is None or user.role not in (
                'organization_admin', 'family_admin', 'fundraiser'):
            _app.abort(403)
        return user

    def supporter_directory_families():
        user = require_supporter_directory_access()
        statement = select(_app.Family).order_by(_app.Family.name)
        if user is not None and user.role != 'organization_admin':
            statement = statement.where(_app.Family.id.in_(select(
                _app.FamilyAssignment.family_id).where(
                    _app.FamilyAssignment.staff_user_id == user.id)))
        return _app.db.session.scalars(statement).all()

    @app.route('/people/new', methods=['GET', 'POST'])
    def new_directory_person():
        require_supporter_directory_access()
        return_to = _app.request.values.get('next', '').strip()
        if not return_to.startswith('/') or return_to.startswith('//'):
            return_to = _app.url_for('supporter_directory')
        if _app.request.method == 'POST':
            name = _app.request.form.get('name', '').strip()
            phone = _app.request.form.get('phone', '').strip()
            email = _app.request.form.get('email', '').strip().lower()
            if not name:
                _app.flash('Enter the person’s name.', 'error')
                return _app.render_template('person_new.html', title='Add person',
                                            return_to=return_to), 400
            if email and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
                _app.flash('Enter a valid email address.', 'error')
                return _app.render_template('person_new.html', title='Add person',
                                            return_to=return_to), 400
            profile = create_neutral_directory_person(name=name, phone=phone, email=email)
            _app.db.session.flush()
            person = canonical_person_for_profile(profile)
            for field_name, limit in (
                    ('home_phone', 80), ('cell_phone', 80),
                    ('home_address', 240), ('city', 120), ('state', 80),
                    ('zip_code', 20), ('workplace', 160),
                    ('work_phone', 80), ('notes', 5000)):
                setattr(person, field_name, _app.request.form.get(
                    field_name, '').strip()[:limit])
            if not person.cell_phone:
                person.cell_phone = phone[:80]
            _app.db.session.commit()
            _app.flash('Person added to the shared name list.')
            return _app.redirect(return_to)
        return _app.render_template('person_new.html', title='Add person',
                                    return_to=return_to)

    @app.route('/supporter-directory', methods=['GET', 'POST'])
    def supporter_directory():
        families = supporter_directory_families()
        import_result = None
        if _app.request.method == 'POST':
            upload = _app.request.files.get('file')
            if upload is None or not upload.filename:
                _app.abort(400, 'Choose a CSV or Excel file.')
            try:
                rows = imported_contact_rows(upload)
            except ValueError as exc:
                _app.abort(400, str(exc))
            created = updated = duplicates = skipped = 0
            seen = set()
            errors = []
            candidates = []
            for row in rows:
                phone = normalized_profile_phone(row['phone'])
                if not row['name'] or not phone:
                    skipped += 1
                    errors.append(f"Row {row['row']}: name and a valid phone number are required.")
                    continue
                if phone in seen:
                    duplicates += 1
                    errors.append(f"Row {row['row']}: duplicate phone number in this file.")
                    continue
                seen.add(phone)
                candidates.append((row, phone))

            # Fetch every possible duplicate at once.  Looking up one profile
            # per spreadsheet row made production imports exceed Gunicorn's
            # request timeout on larger contact lists.
            existing_profiles = {
                profile.normalized_phone: profile for profile in
                _app.db.session.scalars(select(SupporterProfile).where(
                    SupporterProfile.normalized_phone.in_(
                        [phone for _, phone in candidates]))).all()
            } if candidates else {}
            for row, phone in candidates:
                profile = existing_profiles.get(phone)
                if profile:
                    duplicates += 1
                    # The upload can safely fill blanks, but never silently
                    # overwrite established profile data.
                    changed = False
                    if not profile.email and row['email']:
                        profile.email = row['email'][:254]
                        changed = True
                    if not profile.name and row['name']:
                        profile.name = row['name'][:160]
                        changed = True
                    updated += int(changed)
                else:
                    _app.db.session.add(SupporterProfile(
                        name=row['name'][:160], phone=row['phone'][:80],
                        normalized_phone=phone, email=row['email'][:254]))
                    created += 1
            _app.db.session.commit()
            import_result = dict(created=created, updated=updated, duplicates=duplicates,
                                 skipped=skipped, errors=errors[:20])

        query = _app.request.args.get('q', '').strip()[:160]
        statement = select(SupporterProfile).order_by(SupporterProfile.name)
        if query:
            statement = statement.where(_app.or_(
                SupporterProfile.name.icontains(query, autoescape=True),
                SupporterProfile.phone.icontains(query, autoescape=True),
                SupporterProfile.email.icontains(query, autoescape=True)))
        profiles = _app.db.session.scalars(statement).all()
        case_counts = dict(_app.db.session.execute(select(
            _app.Contact.supporter_key, _app.func.count(_app.Contact.id)
        ).where(_app.Contact.supporter_key.in_([
            'phone:' + row.normalized_phone for row in profiles
        ])).group_by(_app.Contact.supporter_key)).all()) if profiles else {}
        return _app.render_template(
            'supporter_directory.html', title='People import', profiles=profiles,
            families=families, import_result=import_result, query=query,
            case_counts=case_counts)

    @app.post('/supporter-directory/<int:profile_id>/connect')
    def connect_supporter_profile(profile_id):
        user = require_supporter_directory_access()
        family_id = _app.request.form.get('family_id', type=int)
        allowed = family_id is not None and (
            user is None or user.role == 'organization_admin' or
            _app.db.session.scalar(select(_app.FamilyAssignment.id).where(
                _app.FamilyAssignment.staff_user_id == user.id,
                _app.FamilyAssignment.family_id == family_id)) is not None)
        if not allowed:
            _app.abort(403, 'Choose a case you can access.')
        profile = _app.db.get_or_404(SupporterProfile, profile_id)
        relationship = _app.request.form.get('relationship', 'Other').strip()
        if relationship not in set(_app.RELATIONSHIPS) | _app.LEGACY_RELATIONSHIPS:
            _app.abort(400, 'Choose a valid relationship.')
        key = 'phone:' + profile.normalized_phone
        duplicate = _app.db.session.scalar(select(_app.Contact.id).where(
            _app.Contact.family_id == family_id,
            _app.Contact.supporter_key == key))
        if duplicate:
            _app.flash('This person is already connected to that case.')
            return _app.redirect(_app.url_for('supporter_directory'))
        existing = _app.db.session.scalar(select(_app.Contact).where(
            _app.Contact.supporter_key == key).order_by(_app.Contact.id))
        contact = _app.Contact(
            family_id=family_id, name=profile.name, phone=profile.phone,
            cell_phone=profile.phone, email=profile.email,
            relationship=relationship, supporter_key=key,
            monthly_cents=existing.monthly_cents if existing else 0,
            pledge_frequency=existing.pledge_frequency if existing else 'Monthly',
            status=existing.status if existing else 'To contact')
        _app.db.session.add(contact)
        _app.db.session.flush()
        attach_supporter_person(contact, source=existing)

        # The case's current Circle of Support is backed by both Contact and
        # SupporterLink.  Imports used to create only the legacy Contact row,
        # which left the person invisible to workflow-aware/fundraiser views.
        link_model = app.extensions['workflows']['models']['SupporterLink']
        user = _app.db.session.get(
            _app.StaffUser, _app.session.get('user_id')) if _app.session.get('user_id') else None
        _app.db.session.add(link_model(
            contact_id=contact.id,
            side='Community',
            relationship=relationship,
            assigned_to=user.id if user and user.role == 'fundraiser' else None,
            permission='Not requested',
            preference='',
            verified=False,
        ))
        sync_followup = app.extensions.get('sync_supporter_followup_task')
        if sync_followup:
            sync_followup(contact)
        _app.db.session.commit()
        _app.flash('Person connected to the case. You can now complete the profile.')
        return _app.redirect(_app.url_for('edit_contact', contact_id=contact.id))


    @app.route('/supporter-directory/<int:profile_id>/edit', methods=['GET', 'POST'])
    def edit_supporter_profile(profile_id):
        require_supporter_directory_access()
        profile = _app.db.get_or_404(SupporterProfile, profile_id)
        person = canonical_person_for_profile(profile)
        _app.db.session.commit()

        def edit_context():
            # Relationship choices include every imported person. Link newly
            # imported profiles in one batch instead of issuing queries and a
            # flush for every row as the directory grows.
            canonicalize_unlinked_profiles()
            _app.db.session.commit()
            relationship_rows = _app.db.session.scalars(select(PersonRelationship).where(
                _app.or_(PersonRelationship.person_one_id == person.id,
                         PersonRelationship.person_two_id == person.id)
            ).order_by(PersonRelationship.id)).all()
            other_ids = {
                row.person_two_id if row.person_one_id == person.id else row.person_one_id
                for row in relationship_rows
            }
            people_by_id = {
                row.id: row for row in _app.db.session.scalars(select(
                    SupporterPerson).where(SupporterPerson.id.in_(other_ids))).all()
            } if other_ids else {}
            related_people = [(
                row, people_by_id.get(
                    row.person_two_id if row.person_one_id == person.id
                    else row.person_one_id))
                for row in relationship_rows]
            connected_ids = {other.id for _, other in related_people if other}
            available_statement = select(SupporterPerson).where(
                SupporterPerson.id != person.id)
            if connected_ids:
                available_statement = available_statement.where(
                    SupporterPerson.id.not_in(connected_ids))
            affiliations = _app.db.session.scalars(select(
                _app.PersonAffiliation).where(
                    _app.PersonAffiliation.person_type == 'supporter_profile',
                    _app.PersonAffiliation.person_id == profile.id
                ).order_by(_app.PersonAffiliation.id)).all()
            used_institutions = {row.institution_id for row in affiliations}
            institution_statement = select(_app.Institution)
            if used_institutions:
                institution_statement = institution_statement.where(
                    _app.Institution.id.not_in(used_institutions))
            return dict(
                profile=profile, person=person,
                person_relationships=related_people,
                relationship_types=PERSON_RELATIONSHIPS,
                available_people=_app.db.session.scalars(
                    available_statement.order_by(SupporterPerson.name)).all(),
                affiliations=affiliations,
                institutions=_app.db.session.scalars(
                    institution_statement.order_by(
                        _app.Institution.kind, _app.Institution.name)).all())
        if _app.request.method == 'POST':
            name = _app.request.form.get('name', '').strip()
            phone = _app.request.form.get('phone', '').strip()
            email = _app.request.form.get('email', '').strip()
            normalized = normalized_profile_phone(phone)
            if not name or (phone and not normalized):
                _app.flash('A name and a valid phone number are required.', 'error')
                return _app.render_template(
                    'supporter_profile_edit.html', title='Edit imported person',
                    **edit_context()), 400

            duplicate = (_app.db.session.scalar(select(SupporterProfile.id).where(
                SupporterProfile.normalized_phone == normalized,
                SupporterProfile.id != profile.id)) if normalized else None)
            old_key = 'phone:' + profile.normalized_phone
            new_key = ('phone:' + normalized) if normalized else person.identity_key
            case_duplicate = None
            if new_key != old_key:
                case_duplicate = _app.db.session.scalar(select(_app.Contact.id).where(
                    _app.Contact.supporter_key == new_key))
            if duplicate or case_duplicate:
                _app.flash('That phone number already belongs to another person.', 'error')
                return _app.render_template(
                    'supporter_profile_edit.html', title='Edit imported person',
                    **edit_context()), 409

            linked_contacts = _app.db.session.scalars(select(_app.Contact).where(
                _app.or_(_app.Contact.person_id == person.id,
                         _app.Contact.supporter_key == old_key))).all()
            profile.name = name[:160]
            profile.phone = phone[:80]
            if normalized:
                profile.normalized_phone = normalized
            profile.email = email[:254]
            values = {
                'name': profile.name, 'phone': profile.phone,
                'email': profile.email, 'supporter_key': new_key,
            }
            for field_name, limit in (
                    ('home_phone', 80), ('cell_phone', 80),
                    ('home_address', 240), ('city', 120), ('state', 80),
                    ('zip_code', 20), ('workplace', 160),
                    ('work_phone', 80), ('notes', 5000)):
                values[field_name] = _app.request.form.get(
                    field_name, '').strip()[:limit]
            if not values['cell_phone']:
                values['cell_phone'] = profile.phone
            try:
                if linked_contacts:
                    person = update_supporter_person(linked_contacts[0], values)
                else:
                    person.identity_key = new_key
                    for field_name in personal_fields:
                        if field_name in values:
                            setattr(person, field_name, values[field_name])
                    sync_person_snapshots(person)
            except ValueError as exc:
                _app.db.session.rollback()
                _app.flash(str(exc), 'error')
                return _app.render_template(
                    'supporter_profile_edit.html', title='Edit imported person',
                    **edit_context()), 409
            _app.db.session.commit()
            _app.flash('Person updated everywhere they are connected.')
            return _app.redirect(_app.url_for('supporter_directory'))

        return _app.render_template(
            'supporter_profile_edit.html', title='Edit imported person',
            **edit_context())

    @app.post('/supporter-directory/<int:profile_id>/relationships')
    def add_person_relationship(profile_id):
        require_supporter_directory_access()
        profile = _app.db.get_or_404(SupporterProfile, profile_id)
        person = canonical_person_for_profile(profile)
        other_id = _app.request.form.get('other_person_id', type=int)
        other = _app.db.session.get(SupporterPerson, other_id)
        relationship = _app.request.form.get('relationship', '').strip()
        if other is None or other.id == person.id:
            _app.abort(400, 'Choose another person.')
        if relationship not in PERSON_RELATIONSHIPS:
            _app.abort(400, 'Choose a valid relationship.')
        one_id, two_id = sorted((person.id, other.id))
        existing = _app.db.session.scalar(select(PersonRelationship.id).where(
            PersonRelationship.person_one_id == one_id,
            PersonRelationship.person_two_id == two_id))
        if existing:
            _app.abort(409, 'These people are already connected.')
        _app.db.session.add(PersonRelationship(
            person_one_id=one_id, person_two_id=two_id,
            relationship=relationship,
            notes=_app.request.form.get('relationship_notes', '').strip()[:500]))
        _app.db.session.commit()
        _app.flash('People connected.')
        return _app.redirect(_app.url_for('edit_supporter_profile', profile_id=profile.id))

    @app.post('/supporter-directory/<int:profile_id>/relationships/<int:relationship_id>/delete')
    def delete_person_relationship(profile_id, relationship_id):
        require_supporter_directory_access()
        profile = _app.db.get_or_404(SupporterProfile, profile_id)
        person = canonical_person_for_profile(profile)
        row = _app.db.get_or_404(PersonRelationship, relationship_id)
        if person.id not in (row.person_one_id, row.person_two_id):
            _app.abort(403)
        _app.db.session.delete(row)
        _app.db.session.commit()
        _app.flash('Person connection removed.')
        return _app.redirect(_app.url_for('edit_supporter_profile', profile_id=profile.id))

    @app.post('/supporter-directory/<int:profile_id>/affiliations')
    def add_profile_affiliation(profile_id):
        require_supporter_directory_access()
        profile = _app.db.get_or_404(SupporterProfile, profile_id)
        institution = _app.db.get_or_404(
            _app.Institution, _app.request.form.get('institution_id', type=int))
        duplicate = _app.db.session.scalar(select(_app.PersonAffiliation.id).where(
            _app.PersonAffiliation.institution_id == institution.id,
            _app.PersonAffiliation.person_type == 'supporter_profile',
            _app.PersonAffiliation.person_id == profile.id))
        if duplicate:
            _app.abort(409, 'This person is already connected to that institution.')
        year_from = _app.request.form.get('year_from', type=int)
        year_to = _app.request.form.get('year_to', type=int)
        if institution.kind == 'Yeshivah' and (
                year_from is None or year_to is None or year_from > year_to):
            _app.abort(400, 'Enter valid attendance years for the yeshivah.')
        _app.db.session.add(_app.PersonAffiliation(
            institution_id=institution.id, person_type='supporter_profile',
            person_id=profile.id,
            grade=_app.request.form.get('grade', '').strip()[:80],
            year_from=year_from, year_to=year_to,
            note=_app.request.form.get('affiliation_note', '').strip()[:300]))
        _app.db.session.commit()
        _app.flash('Institution connected to the person.')
        return _app.redirect(_app.url_for('edit_supporter_profile', profile_id=profile.id))

    @app.post('/supporter-directory/<int:profile_id>/affiliations/<int:affiliation_id>/delete')
    def delete_profile_affiliation(profile_id, affiliation_id):
        require_supporter_directory_access()
        profile = _app.db.get_or_404(SupporterProfile, profile_id)
        row = _app.db.get_or_404(_app.PersonAffiliation, affiliation_id)
        if row.person_type != 'supporter_profile' or row.person_id != profile.id:
            _app.abort(403)
        _app.db.session.delete(row)
        _app.db.session.commit()
        _app.flash('Institution connection removed.')
        return _app.redirect(_app.url_for('edit_supporter_profile', profile_id=profile.id))

    def add_shared_gabbai(family_id):
        user = (_app.db.session.get(_app.StaffUser, _app.session.get('user_id'))
                if _app.session.get('user_id') else None)
        if not app.config['DEMO']:
            if user is None or user.role not in (
                    'organization_admin', 'family_admin', 'office_employee'):
                _app.abort(403)
            if user.role != 'organization_admin':
                assigned = _app.db.session.scalar(select(
                    _app.FamilyAssignment.id).where(
                        _app.FamilyAssignment.staff_user_id == user.id,
                        _app.FamilyAssignment.family_id == family_id))
                if assigned is None:
                    _app.abort(403)
        family = _app.db.get_or_404(_app.Family, family_id)
        shuls = _family_profile_shuls(family)
        institution_id = _app.request.form.get('institution_id', type=int)
        institution = next((row for row in shuls if row.id == institution_id), None)
        if institution is None and len(shuls) == 1:
            institution = shuls[0]
        if institution is None:
            legacy_name = (family.weekday_shul or family.shabbos_shul or '').strip()
            if legacy_name:
                institution = _find_shul(legacy_name)
                if institution is None:
                    institution = _app.Institution(kind='Shul', name=legacy_name)
                    _app.db.session.add(institution)
                    _app.db.session.flush()
                linked = _app.db.session.scalar(select(_app.PersonAffiliation.id).where(
                    _app.PersonAffiliation.institution_id == institution.id,
                    _app.PersonAffiliation.person_type == 'family',
                    _app.PersonAffiliation.person_id == family.id))
                if linked is None:
                    _app.db.session.add(_app.PersonAffiliation(
                        institution_id=institution.id, person_type='family',
                        person_id=family.id))
        if institution is None:
            _app.abort(400, 'Choose which shul this gabbai belongs to.')
        name = (_app.request.form.get('name') or '').strip()
        phone = (_app.request.form.get('phone') or '').strip()
        if not name:
            _app.abort(400, 'Gabbai name is required.')
        if len(name) > 160 or len(phone) > 80:
            _app.abort(400, 'Value is too long.')
        person = _canonical_helper(name, [phone])
        if person is not None:
            _attach_helper(institution, person, 'shul_gabbai')
        sync_shul_gabbai_supporters()
        actor = user.email if user else 'Demo user'
        _app.db.session.add(_app.Audit(
            actor=actor, action='Added shared shul gabbai', family_id=family.id))
        _app.db.session.commit()
        _app.flash('Shul gabbai added.')
        return _app.redirect(_app.url_for('family_detail', family_id=family.id))

    app.view_functions['add_gabbai'] = add_shared_gabbai

    @app.post('/community-helpers/<int:helper_id>')
    def update_shared_helper(helper_id):
        user = (_app.db.session.get(_app.StaffUser, _app.session.get('user_id'))
                if _app.session.get('user_id') else None)
        if not app.config['DEMO'] and (user is None or user.role != 'organization_admin'):
            _app.abort(403)
        helper = _app.db.get_or_404(HelperPerson, helper_id)
        name = ' '.join((_app.request.form.get('name') or '').split())[:160]
        phone = (_app.request.form.get('phone') or '').strip()[:80]
        if not name:
            _app.abort(400, 'Gabbai name is required.')
        normalized = _normalize_rabbi_name(name)
        duplicate = _app.db.session.scalar(select(HelperPerson.id).where(
            HelperPerson.normalized_name == normalized,
            HelperPerson.id != helper.id))
        if duplicate:
            _app.abort(409, 'A helper with this name already exists.')
        helper.name = name
        helper.normalized_name = normalized
        for saved_phone in _app.db.session.scalars(select(HelperPhone).where(
                HelperPhone.helper_person_id == helper.id)).all():
            _app.db.session.delete(saved_phone)
        _app.db.session.flush()
        if phone:
            _app.db.session.add(HelperPhone(helper_person_id=helper.id, phone=phone))
        _app.db.session.flush()
        sync_shul_gabbai_supporters()
        _app.db.session.add(_app.Audit(
            actor=user.email if user else 'Demo user',
            action='Updated shared shul helper'))
        _app.db.session.commit()
        _app.flash('Shul gabbai updated.')
        return _app.redirect(_app.request.referrer or _app.url_for(
            'community_directories', kind='Shul'))

    @app.get('/api/shul-rabbis')
    def shul_rabbis_api():
        rows = _app.db.session.execute(select(_app.Institution, ShulRabbi).join(
            ShulRabbi, ShulRabbi.institution_id == _app.Institution.id, isouter=True).where(
            _app.Institution.kind == 'Shul')).all()
        result = {}
        for institution, assignment in rows:
            phones = _rabbi_phones(institution) if assignment else []
            result[institution.name] = {
                'name': assignment.rabbi_name if assignment else '',
                'phone': phones[0] if phones else '',
                'phones': phones,
                'assistants': _assistant_payload(institution),
            }
        return result

    @app.get('/api/shul-gabbais')
    def shul_gabbais_api():
        shuls = _app.db.session.scalars(select(_app.Institution).where(
            _app.Institution.kind == 'Shul').order_by(_app.Institution.name)).all()
        result = {}
        for shul in shuls:
            result[shul.name] = _helper_payloads(shul, 'shul_gabbai')
        return result

    @app.context_processor
    def shul_rabbi_context():
        # This directory data is only rendered on the family and community
        # directory screens. Loading it globally caused an N+1 query storm on
        # every page (dashboard, lists, payouts, etc.): each shul triggered
        # separate phone, assistant, gabbai and gabbai-phone queries.
        #
        # Keep harmless defaults available to every template, but do no
        # directory work unless the current screen can actually use it.
        empty_context = {
            'rabbi_people': [],
            'family_rabbi_preference': lambda family_id: None,
            'family_gabbai_contacts': lambda family_id: [],
            'family_shul_choices': lambda family: [],
            'family_rabbi_assistant_contacts': lambda family: [],
            'family_yeshivah_history': lambda family_id: [],
            'shul_rabbi_map': {},
            'family_phone_values': lambda family: (
                _clean_phones(_app.request.form.getlist('phone'))
                if _app.request.method == 'POST' else _family_phones(family)),
        }
        if _app.request.endpoint not in {
                'new_family', 'edit_family', 'family_detail', 'community_directories'}:
            return empty_context

        # A family profile can display contacts for only its weekday and
        # Shabbos shuls.  Do not build the entire organization-wide directory
        # for that page: doing so used to issue several queries per shul and
        # made profiles progressively slower as the directory grew.
        institution_filter = [_app.Institution.kind == 'Shul']
        profile_institution_ids = []
        if _app.request.endpoint == 'family_detail':
            family_id = (_app.request.view_args or {}).get('family_id')
            # family_detail already loaded this object into the request's
            # identity map, so this does not require another database query.
            context_family = _app.db.session.get(_app.Family, family_id)
            profile_institution_ids = list(_app.db.session.scalars(select(
                _app.PersonAffiliation.institution_id
            ).join(_app.Institution).where(
                _app.PersonAffiliation.person_type == 'family',
                _app.PersonAffiliation.person_id == family_id,
                _app.Institution.kind == 'Shul',
                _app.PersonAffiliation.note.contains('Family profile'),
            )).all())
            relevant_names = {name for name in (
                context_family.weekday_shul if context_family else '',
                context_family.shabbos_shul if context_family else '') if name}
            if profile_institution_ids:
                institution_filter.append(_app.Institution.id.in_(profile_institution_ids))
                rows = _app.db.session.execute(select(_app.Institution, ShulRabbi).join(
                    ShulRabbi, ShulRabbi.institution_id == _app.Institution.id,
                    isouter=True).where(*institution_filter)).all()
            elif not relevant_names:
                rows = []
            else:
                institution_filter.append(_app.Institution.name.in_(relevant_names))
                rows = _app.db.session.execute(select(_app.Institution, ShulRabbi).join(
                    ShulRabbi, ShulRabbi.institution_id == _app.Institution.id,
                    isouter=True).where(*institution_filter)).all()
        else:
            # The intake form gets its shul names separately and the directory
            # screen renders its own selected records; neither consumes this
            # contact map.
            rows = []
        mapping = {}
        institution_contacts = {}
        for institution, assignment in rows:
            phones = _rabbi_phones(institution) if assignment else []
            gabbais = _helper_payloads(institution, 'shul_gabbai')
            assistants = _assistant_payload(institution)
            if _app.request.endpoint == 'family_detail':
                assistant_role = {
                    'yi': 'גבאי פונעם רב',
                    'he': 'גבאי הרב',
                }.get(_app.session.get('language', 'en'), 'Rabbi assistant / gabbai')
                assistants = [
                    {**assistant, 'name': f"{assistant['name']} · {assistant_role}"}
                    for assistant in assistants
                ]
            payload = {
                'name': assignment.rabbi_name if assignment else '',
                'phone': phones[0] if phones else '',
                'phones': phones,
                'assistants': assistants,
                'gabbais': gabbais,
            }
            institution_contacts[institution.id] = payload
            mapping[institution.name] = payload
        people = (_app.db.session.scalars(select(RabbiPerson).order_by(RabbiPerson.name)).all()
                  if _app.request.endpoint in {'new_family', 'edit_family', 'community_directories'}
                  else [])
        def family_rabbi_preference(family_id):
            return _app.db.session.get(FamilyRabbiPreference, family_id) if family_id else None
        def family_gabbai_contacts(family_id):
            if not family_id:
                return []
            family = _app.db.session.get(_app.Family, family_id)
            result = []
            seen = set()

            # Gabbaim belong to the shared shul, not to an individual family.
            # Read the current shul directory first so applicants who selected
            # the same shul also see gabbaim added after their profile was saved.
            linked_ids = profile_institution_ids if family_id == (
                (_app.request.view_args or {}).get('family_id')) else list(
                    _app.db.session.scalars(select(
                        _app.PersonAffiliation.institution_id
                    ).join(_app.Institution).where(
                        _app.PersonAffiliation.person_type == 'family',
                        _app.PersonAffiliation.person_id == family_id,
                        _app.Institution.kind == 'Shul',
                        _app.PersonAffiliation.note.contains('Family profile'),
                    )).all())
            contact_groups = [institution_contacts.get(row_id, {}) for row_id in linked_ids]
            if not contact_groups:
                contact_groups = [mapping.get(name, {}) for name in (
                    (family.weekday_shul, family.shabbos_shul) if family is not None else ())]
            for contacts in contact_groups:
                for gabbai in contacts.get('gabbais', []):
                    identity = (gabbai['name'].casefold(), tuple(gabbai['phones']))
                    if identity in seen:
                        continue
                    seen.add(identity)
                    result.append(gabbai)

            return result
        def family_rabbi_assistant_contacts(family):
            if family is None:
                return []
            result = []
            seen = set()
            for shul_name in (family.weekday_shul, family.shabbos_shul):
                contacts = mapping.get(shul_name, {})
                if (_normalize_rabbi_name(contacts.get('name')) !=
                        _normalize_rabbi_name(family.rabbi)):
                    continue
                for assistant in contacts.get('assistants', []):
                    identity = assistant['name'].casefold()
                    if identity in seen:
                        continue
                    seen.add(identity)
                    result.append(assistant)
            return result
        def family_yeshivah_history(family_id):
            return _app.db.session.execute(select(_app.PersonAffiliation, _app.Institution).join(
                _app.Institution, _app.Institution.id == _app.PersonAffiliation.institution_id
            ).where(
                _app.PersonAffiliation.person_type == 'family',
                _app.PersonAffiliation.person_id == family_id,
                _app.Institution.kind == 'Yeshivah',
            ).order_by(_app.PersonAffiliation.year_from, _app.PersonAffiliation.id)).all()
        return {
            'rabbi_people': people,
            'family_rabbi_preference': family_rabbi_preference,
            'family_gabbai_contacts': family_gabbai_contacts,
            'family_shul_choices': _family_profile_shuls,
            'family_rabbi_assistant_contacts': family_rabbi_assistant_contacts,
            'family_yeshivah_history': family_yeshivah_history,
            'shul_rabbi_map': mapping,
            'family_phone_values': lambda family: (
                _clean_phones(_app.request.form.getlist('phone'))
                if _app.request.method == 'POST' else _family_phones(family)),
        }

    def require_org_admin():
        if app.config['DEMO']:
            return
        user = _app.db.session.get(_app.StaffUser, _app.session.get('user_id'))
        if user is None or user.role != 'organization_admin':
            _app.abort(403)

    def add_audit(action):
        user_id = _app.session.get('user_id') if has_request_context() else None
        user = (_app.db.session.get(_app.StaffUser, user_id) if user_id else None)
        _app.db.session.add(_app.Audit(
            actor=user.email if user else 'System', action=action))

    @app.post('/community-directories/shuls/<int:institution_id>/rabbis')
    def add_shul_rabbi(institution_id):
        require_org_admin()
        institution = _app.db.get_or_404(_app.Institution, institution_id)
        if institution.kind != 'Shul':
            _app.abort(400, 'Choose a valid shul.')
        person_id = _app.request.form.get('rabbi_person_id', type=int)
        if person_id:
            person = _app.db.get_or_404(RabbiPerson, person_id)
        else:
            name = _app.request.form.get('rabbi_name', '').strip()[:160]
            if not name:
                _app.abort(400, 'Rabbi name is required.')
            person = _canonical_rabbi(
                name, _app.request.form.get('rabbi_phone', '').strip()[:80])
        association = _attach_rabbi(
            institution, person, make_primary=_app.request.form.get('make_primary') == '1')
        _sync_automatic_family_rabbis()
        add_audit(f'Connected rabbi to shul: {person.name} / {institution.name}')
        _app.db.session.commit()
        _app.flash('Rabbi connected.')
        return _app.redirect(_app.url_for('community_directories', kind='Shul',
                                          network_id=institution.id))

    @app.post('/community-directories/shul-rabbis/<int:association_id>/primary')
    def make_shul_rabbi_primary(association_id):
        require_org_admin()
        association = _app.db.get_or_404(ShulRabbiAssociation, association_id)
        _attach_rabbi(association.institution, association.rabbi_person, make_primary=True)
        _sync_automatic_family_rabbis()
        add_audit(f'Changed primary rabbi: {association.institution.name}')
        _app.db.session.commit()
        return _app.redirect(_app.url_for('community_directories', kind='Shul',
                                          network_id=association.institution_id))

    @app.post('/community-directories/shul-rabbis/<int:association_id>/delete')
    def delete_shul_rabbi(association_id):
        require_org_admin()
        association = _app.db.get_or_404(ShulRabbiAssociation, association_id)
        institution = association.institution
        was_primary = association.is_primary
        name = association.rabbi_person.name
        _app.db.session.delete(association)
        _app.db.session.flush()
        if was_primary:
            replacement = _app.db.session.scalar(select(ShulRabbiAssociation).where(
                ShulRabbiAssociation.institution_id == institution.id
            ).order_by(ShulRabbiAssociation.id))
            if replacement:
                replacement.is_primary = True
                _app.db.session.flush()
        _sync_legacy_primary(institution.id)
        _sync_automatic_family_rabbis()
        add_audit(f'Removed rabbi from shul: {name} / {institution.name}')
        _app.db.session.commit()
        return _app.redirect(_app.url_for('community_directories', kind='Shul',
                                          network_id=institution.id))

    for endpoint in ('new_family', 'edit_family'):
        original = app.view_functions.get(endpoint)
        if original is None:
            continue

        def wrapped(*args, __original=original, __endpoint=endpoint, **kwargs):
            if (_app.request.method == 'POST' and
                    _app.request.form.get('rabbi_mode') == 'canonical'):
                selected_id = _app.request.form.get('rabbi_person_id', type=int)
                if selected_id is None or _app.db.session.get(RabbiPerson, selected_id) is None:
                    _app.abort(400, 'Choose a valid rabbi.')
            response = app.make_response(__original(*args, **kwargs))
            if _app.request.method == 'POST' and response.status_code < 400:
                family_id = kwargs.get('family_id') if __endpoint == 'edit_family' else _family_id_from_response(response)
                if family_id:
                    # The original handler has already committed the family
                    # record and its selected institution affiliations. Old
                    # production directory rows can still conflict while the
                    # compatibility projections are synchronized. That
                    # secondary conflict must not turn a successful profile
                    # save into a 409 page or discard the selected shul.
                    sync_error = None
                    # A concurrent edit can invalidate the compatibility
                    # projection after the family itself has committed.  A
                    # rollback clears that failed projection transaction, so
                    # retry it once from the newly committed family record
                    # instead of reporting success while silently dropping
                    # the rabbi/gabbai details.
                    for attempt in range(2):
                        try:
                            _save_shul_connections(int(family_id))
                            sync_shul_gabbai_supporters([int(family_id)])
                            _app.db.session.commit()
                            sync_error = None
                            break
                        except (IntegrityError, StaleDataError) as exc:
                            sync_error = exc
                            _app.db.session.rollback()
                            if attempt == 0:
                                continue
                    if sync_error is not None:
                        app.logger.error(
                            'Family %s saved, but its shul directory sync failed twice',
                            family_id, exc_info=(type(sync_error), sync_error,
                                                 sync_error.__traceback__))
                    _replace_family_phones(
                        int(family_id), _app.request.form.getlist('phone'))
                    _app.db.session.commit()
            return response

        app.view_functions[endpoint] = wrapped

    TASK_STATUSES = ('To do', 'In progress', 'Waiting', 'Completed', 'Cancelled')
    TASK_OUTREACH_RESULTS = (
        'No answer', 'Left a message', 'Contacted', 'Wrong number',
        'Paused', 'Declined')
    TASK_PRIORITIES = ('Low', 'Normal', 'High', 'Urgent')

    def task_user():
        user_id = _app.session.get('user_id')
        return _app.db.session.get(_app.StaffUser, user_id) if user_id else None

    def task_is_admin(user):
        return bool(user and user.role == 'organization_admin')

    def communication_contact(contact_id):
        contact = _app.db.get_or_404(_app.Contact, contact_id)
        user = task_user()
        if task_is_admin(user):
            return contact
        assigned = user and _app.db.session.scalar(select(_app.FamilyAssignment.id).where(
            _app.FamilyAssignment.staff_user_id == user.id,
            _app.FamilyAssignment.family_id == contact.family_id))
        if not assigned:
            _app.abort(403, 'You are not assigned to this family.')
        if user.role == 'fundraiser' and app.extensions['workflows']['enforced']():
            link_model = app.extensions['workflows']['models']['SupporterLink']
            link = _app.db.session.get(link_model, contact.id)
            if not link or link.assigned_to != user.id:
                _app.abort(403, 'You are not assigned to this supporter.')
        return contact

    def communication_row(contact, kind, subject='', body='', status='completed',
                          scheduled_for=None, email_message=None,
                          provider_message_id='', delivery_error='',
                          direction='outbound'):
        now = _app.datetime.now(_app.timezone.utc).replace(tzinfo=None)
        row = SupporterCommunication(
            contact_id=contact.id, family_id=contact.family_id,
            staff_user_id=task_user().id if task_user() else None,
            email_message_id=email_message.id if email_message else None,
            kind=kind, direction=direction, subject=subject, body=body, status=status,
            provider_message_id=provider_message_id or '',
            delivery_error=delivery_error or '',
            scheduled_for=scheduled_for,
            completed_at=now if status == 'completed' else None)
        _app.db.session.add(row)
        return row

    def family_for_inbound_phone(sender):
        """Match an inbound text to one unambiguous applicant phone number."""
        try:
            sender = normalize_phone(sender.replace('whatsapp:', '', 1))
        except ValueError:
            return None
        matches = []
        for family in _app.db.session.scalars(select(_app.Family)):
            for value in _family_phones(family):
                try:
                    if normalize_phone(value) == sender:
                        matches.append(family)
                        break
                except ValueError:
                    continue
        return matches[0] if len(matches) == 1 else None

    def contact_for_inbound_phone(sender, channel):
        """Match a reply to the supporter most recently contacted on this channel."""
        try:
            sender = normalize_phone(sender.replace('whatsapp:', '', 1))
        except ValueError:
            return None
        matches = []
        for contact in _app.db.session.scalars(select(_app.Contact)):
            for value in (contact.cell_phone, contact.phone, contact.home_phone,
                          contact.work_phone):
                if not value:
                    continue
                try:
                    if normalize_phone(value) == sender:
                        matches.append(contact)
                        break
                except ValueError:
                    continue
        if not matches:
            return None
        contact_ids = [contact.id for contact in matches]
        latest = _app.db.session.scalar(select(SupporterCommunication).where(
            SupporterCommunication.contact_id.in_(contact_ids),
            SupporterCommunication.kind == channel,
            SupporterCommunication.direction == 'outbound',
        ).order_by(SupporterCommunication.created_at.desc(),
                   SupporterCommunication.id.desc()))
        if latest:
            return next(contact for contact in matches if contact.id == latest.contact_id)
        return min(matches, key=lambda contact: contact.id)

    def sms_phone_key(value):
        try:
            return normalize_phone((value or '').replace('whatsapp:', '', 1))
        except ValueError:
            return ''

    def sms_person_identity(phone):
        """Resolve one number without guessing when distinct people share it."""
        normalized = sms_phone_key(phone)
        cache = getattr(_app.g, 'sms_identity_cache', None)
        if cache is None:
            cache = _app.g.sms_identity_cache = {}
        if normalized in cache:
            return cache[normalized]
        unknown = {
            'known': False, 'name': 'Unknown contact', 'url': '',
            'normalized_phone': normalized or (phone or ''), 'ambiguous': False}
        if not normalized:
            return unknown
        manual = _app.db.session.get(SmsPhoneLink, normalized)
        if manual and manual.profile:
            cache[normalized] = {
                'known': True, 'name': manual.profile.name,
                'url': _app.url_for('edit_supporter_profile',
                                    profile_id=manual.profile_id),
                'normalized_phone': normalized, 'ambiguous': False,
                'profile_id': manual.profile_id, 'confirmed': True}
            return cache[normalized]

        candidates = {}

        def add(key, name, url, profile_id=None):
            # The same supporter can appear as a profile, canonical person,
            # and case contact. Those rows are one SMS identity.
            name_key = ' '.join((name or '').split()).casefold()
            for candidate in candidates.values():
                if candidate['name_key'] == name_key:
                    if key[0] == 'contact':
                        candidate['url'] = url
                    return
            if key not in candidates:
                candidates[key] = {
                    'known': True, 'name': name, 'url': url,
                    'normalized_phone': normalized, 'ambiguous': False,
                    'profile_id': profile_id, 'confirmed': False,
                    'name_key': name_key}

        directory = getattr(_app.g, 'sms_directory_cache', None)
        if directory is None:
            directory = _app.g.sms_directory_cache = {
                'profiles': _app.db.session.scalars(select(SupporterProfile)).all(),
                'people': _app.db.session.scalars(select(_app.SupporterPerson)).all(),
                'contacts': _app.db.session.scalars(select(_app.Contact)).all(),
                'families': _app.db.session.scalars(select(_app.Family)).all(),
                'staff': _app.db.session.scalars(select(_app.StaffUser)).all(),
                'askanim': _app.db.session.scalars(select(_app.Askan)).all()}
            phone_index = directory['phone_index'] = {
                kind: {} for kind in ('profiles', 'people', 'contacts',
                                     'families', 'staff', 'askanim')}
            fields = {
                'profiles': ('phone',),
                'people': ('phone', 'cell_phone', 'home_phone', 'work_phone'),
                'contacts': ('phone', 'cell_phone', 'home_phone', 'work_phone'),
                'families': ('phone',), 'staff': ('phone',),
                'askanim': ('phone',)}
            for kind, rows in directory.items():
                if kind == 'phone_index':
                    continue
                for row in rows:
                    for field in fields[kind]:
                        key = sms_phone_key(getattr(row, field))
                        if key:
                            phone_index[kind].setdefault(key, {})[row.id] = row
        matches = directory['phone_index']
        profiles = directory['profiles']
        profiles_by_person = {
            row.person_id: row for row in profiles if row.person_id is not None}
        for profile in matches['profiles'].get(normalized, {}).values():
            add(('person', profile.person_id) if profile.person_id else
                ('profile', profile.id), profile.name,
                _app.url_for('edit_supporter_profile', profile_id=profile.id),
                profile.id)
        for person in matches['people'].get(normalized, {}).values():
            profile = profiles_by_person.get(person.id)
            add(('person', person.id), person.name,
                (_app.url_for('edit_supporter_profile', profile_id=profile.id)
                 if profile else _app.url_for('supporter_directory')),
                profile.id if profile else None)
        for contact in matches['contacts'].get(normalized, {}).values():
            add(('person', contact.person_id) if contact.person_id else
                ('contact', contact.id), contact.name,
                _app.url_for('supporter_detail', contact_id=contact.id))
        for family in matches['families'].get(normalized, {}).values():
            add(('family', family.id), family.name,
                _app.url_for('family_detail', family_id=family.id))
        for staff in matches['staff'].get(normalized, {}).values():
            add(('staff', staff.id), staff.name or staff.email,
                _app.url_for('people_access'))
        for askan in matches['askanim'].get(normalized, {}).values():
            add(('askan', askan.id), askan.name,
                _app.url_for('network_askan_detail', askan_id=askan.id))
        if len(candidates) == 1:
            cache[normalized] = next(iter(candidates.values()))
            return cache[normalized]
        if len(candidates) > 1:
            unknown['ambiguous'] = True
            unknown['matches'] = list(candidates.values())
        cache[normalized] = unknown
        return unknown

    def sms_contact_ids(phone):
        normalized = sms_phone_key(phone)
        directory = getattr(_app.g, 'sms_directory_cache', None)
        if directory:
            return list(directory['phone_index']['contacts'].get(normalized, {}))
        return [
            contact.id for contact in (directory['contacts'] if directory else
                _app.db.session.scalars(select(_app.Contact)).all())
            if any(sms_phone_key(value) == normalized for value in (
                contact.phone, contact.cell_phone, contact.home_phone,
                contact.work_phone))
        ]

    def sms_conversation(phone):
        """Combine general and supporter SMS records into one dated thread."""
        normalized = sms_phone_key(phone)
        rows = []
        general_index = getattr(_app.g, 'sms_general_index', None)
        if general_index is None:
            general_index = _app.g.sms_general_index = {}
            for message in _app.db.session.scalars(select(GeneralSmsMessage)):
                general_index.setdefault(sms_phone_key(message.phone), []).append(message)
        for message in general_index.get(normalized, ()):
            rows.append({
                    'direction': message.direction, 'body': message.body,
                    'created_at': message.created_at, 'status': message.status,
                    'family': message.family, 'provider_id': message.provider_message_id,
                    'source_id': f'general-{message.id}'})
        contact_ids = sms_contact_ids(normalized)
        if contact_ids:
            communications = _app.db.session.scalars(select(
                SupporterCommunication).where(
                    SupporterCommunication.contact_id.in_(contact_ids),
                    SupporterCommunication.kind == 'sms')).all()
            for message in communications:
                rows.append({
                    'direction': message.direction, 'body': message.body,
                    'created_at': message.created_at, 'status': message.status,
                    'family': message.family, 'provider_id': message.provider_message_id,
                    'source_id': f'supporter-{message.id}'})
        deduplicated = {}
        for row in rows:
            key = ('provider', row['provider_id']) if row['provider_id'] else (
                row['direction'], row['body'], row['created_at'])
            deduplicated.setdefault(key, row)
        conversation = sorted(deduplicated.values(),
                              key=lambda row: (row['created_at'], row['source_id']))
        campaign_cache = getattr(_app.g, 'sms_campaign_cache', None)
        if campaign_cache is None:
            campaign_cache = _app.g.sms_campaign_cache = {}
        for row in conversation:
            if row['family'] and row['family'].id not in campaign_cache:
                campaign_cache[row['family'].id] = _app.db.session.scalar(select(
                    _app.CharityCampaign).where(
                    _app.CharityCampaign.family_id == row['family'].id).order_by(
                    _app.CharityCampaign.id.desc()))
            row['campaign'] = (campaign_cache.get(row['family'].id)
                               if row['family'] else None)
        return conversation

    def parse_communication_time(value):
        try:
            return _app.datetime.strptime(value, '%Y-%m-%dT%H:%M')
        except (TypeError, ValueError):
            _app.abort(400, 'Enter a valid follow-up date and time.')

    @app.template_filter('callback_time')
    def callback_time(value):
        """Scheduled callbacks are stored as Eastern local wall times."""
        if value is None:
            return ''
        from zoneinfo import ZoneInfo
        return value.replace(tzinfo=ZoneInfo('America/New_York')).strftime(
            '%m/%d/%Y %I:%M %p %Z')

    def supporter_pledge_delivery(contact):
        """Choose the payment destination from the supporter's current pledges.

        Contact rows are case-specific, while ``supporter_key`` is the shared
        billing identity. Recompute this at render/send time so adding a second
        family immediately switches the supporter to one Yazory link.
        """
        linked = _app.db.session.scalars(select(_app.Contact).where(
            _app.Contact.supporter_key == contact.supporter_key
        ).order_by(_app.Contact.family_id, _app.Contact.id)).all() \
            if contact.supporter_key else [contact]
        active = [row for row in linked
                  if row.monthly_cents > 0 and row.status not in ('Paused', 'Declined')]
        by_family = {}
        for row in active:
            by_family.setdefault(row.family_id, row)
        pledges = list(by_family.values()) or [contact]
        if len(pledges) > 1:
            return {
                'kind': 'Yazory',
                'url': _public_url('supporter_donation', contact_id=contact.id),
                'pledges': pledges,
                'reason': f'{len(pledges)} connected family pledges · one charge',
            }
        pledge = pledges[0]
        campaign = _app.db.session.scalar(select(_app.CharityCampaign).where(
            _app.CharityCampaign.family_id == pledge.family_id).order_by(
                _app.CharityCampaign.id.desc()))
        if campaign and campaign.public_url:
            return {'kind': 'ABCharity', 'url': campaign.public_url,
                    'pledges': pledges,
                    'reason': f'One family pledge · campaign {campaign.external_id}'}
        if campaign:
            return {
                'kind': 'ABCharity link missing',
                'url': '',
                'pledges': pledges,
                'family_id': pledge.family_id,
                'reason': 'Add the public ABCharity campaign link before sending this pledge.',
            }
        return {
            'kind': 'Yazory',
            'url': _public_url('supporter_donation', contact_id=contact.id),
            'pledges': pledges,
            'reason': 'ABCharity campaign is unavailable · secure Yazory payment',
        }

    def render_communications(ai_contact=None, ai_subject='', ai_body=''):
        user = task_user()
        if user is None or user.role not in (
                'organization_admin', 'family_admin', 'fundraiser'):
            _app.abort(403)
        selected_contact_id = (ai_contact.id if ai_contact else
                               _app.request.args.get('contact_id', type=int))
        supporter_query = _app.request.args.get('supporter_q', '').strip()[:160]
        supporter_page = max(
            _app.request.args.get('supporter_page', 1, type=int), 1)
        external_email = _app.request.args.get('recipient_email', '').strip().lower()[:254]
        sms_phone = _app.request.args.get('sms_phone', '').strip()[:80]
        sms_name = _app.request.args.get('sms_name', '').strip()[:160]
        sms_family_id = _app.request.args.get('family_id', type=int)
        sms_family = None
        if sms_phone:
            if user.role != 'organization_admin':
                _app.abort(403)
            if sms_family_id is not None:
                sms_family = _app.db.get_or_404(_app.Family, sms_family_id)
        if external_email and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', external_email):
            _app.abort(400, 'Enter a valid email address.')
        # The overview is intentionally bounded.  A specific supporter remains
        # directly addressable from Tasks and Supporters, while avoiding the
        # former all-supporters render (and its per-supporter pledge queries).
        contact_statement = select(_app.Contact).options(
            joinedload(_app.Contact.family),
            joinedload(_app.Contact.parent_supporter),
        ).order_by(_app.Contact.name)
        if selected_contact_id is not None:
            contact_statement = contact_statement.where(
                _app.Contact.id == selected_contact_id)
        elif supporter_query:
            contact_statement = contact_statement.where(_app.or_(
                _app.Contact.name.icontains(supporter_query, autoescape=True),
                _app.Contact.phone.icontains(supporter_query, autoescape=True),
                _app.Contact.cell_phone.icontains(supporter_query, autoescape=True),
                _app.Contact.email.icontains(supporter_query, autoescape=True),
            ))
        if not task_is_admin(user):
            contact_statement = contact_statement.where(_app.Contact.family_id.in_(select(
                _app.FamilyAssignment.family_id).where(
                    _app.FamilyAssignment.staff_user_id == user.id)))
        if user.role == 'fundraiser' and app.extensions['workflows']['enforced']():
            link_model = app.extensions['workflows']['models']['SupporterLink']
            contact_statement = contact_statement.where(_app.Contact.id.in_(select(
                link_model.contact_id).where(link_model.assigned_to == user.id)))
        if selected_contact_id is None:
            contact_statement = contact_statement.offset(
                (supporter_page - 1) * 75).limit(76)
        contacts = _app.db.session.scalars(contact_statement).unique().all()
        supporters_have_next = selected_contact_id is None and len(contacts) > 75
        if supporters_have_next:
            contacts = contacts[:75]
        if selected_contact_id is not None and not contacts:
            _app.abort(404)
        selected_contact = contacts[0] if selected_contact_id is not None else None
        contact_ids = [row.id for row in contacts]
        history = (_app.db.session.scalars(select(SupporterCommunication).where(
            SupporterCommunication.contact_id.in_(contact_ids)).order_by(
                SupporterCommunication.created_at.desc(),
                SupporterCommunication.id.desc()).limit(300)).all()
            if contact_ids else [])
        latest = {}
        for row in history:
            latest.setdefault(row.contact_id, row)
        due = [row for row in history if row.status == 'scheduled']
        from zoneinfo import ZoneInfo
        now = _app.datetime.now(ZoneInfo('America/New_York')).replace(tzinfo=None)
        callback_contact_ids = {row.contact_id for row in due}
        overdue_contact_ids = {
            row.contact_id for row in due
            if row.scheduled_for and row.scheduled_for < now
        }
        email_replies = [row for row in history
                         if row.kind == 'email_reply' and row.status == 'received']
        applicant_history = []
        if selected_contact is None and user.role != 'fundraiser':
            applicant_statement = select(ApplicantMessage).order_by(
                ApplicantMessage.created_at.desc(), ApplicantMessage.id.desc()).limit(300)
            if not task_is_admin(user):
                applicant_statement = applicant_statement.where(
                    ApplicantMessage.family_id.in_(select(
                        _app.FamilyAssignment.family_id).where(
                            _app.FamilyAssignment.staff_user_id == user.id)))
            applicant_history = _app.db.session.scalars(applicant_statement).all()
        applicant_replies = [row for row in applicant_history
                             if row.direction == 'applicant' and row.status == 'unread']
        inbox_history = None
        if selected_contact is None and user.role == 'organization_admin':
            inbox_history = _app.db.session.scalars(select(InboundInboxMessage).order_by(
                InboundInboxMessage.created_at.desc(),
                InboundInboxMessage.id.desc()).limit(300)).all()
        inbox_messages = [row for row in (inbox_history or []) if row.status == 'unread']
        general_sms_history = None
        if selected_contact is None and user.role == 'organization_admin':
            general_sms_history = _app.db.session.scalars(
                select(GeneralSmsMessage).order_by(
                    GeneralSmsMessage.created_at.desc(),
                    GeneralSmsMessage.id.desc()).limit(300)).all()
        general_sms_messages = [
            row for row in (general_sms_history or [])
            if row.direction == 'inbound' and row.status == 'unread'
        ]
        sms_identities = {}
        sms_threads = {}
        sms_context = {}
        if general_sms_history is not None:
            for message in general_sms_history:
                key = sms_phone_key(message.phone) or message.phone
                sms_identities[key] = sms_person_identity(message.phone)
                if key not in sms_threads:
                    sms_threads[key] = sms_conversation(message.phone)
                if message.direction == 'inbound':
                    prior = [
                        row for row in sms_threads[key]
                        if row['direction'] == 'outbound'
                        and row['created_at'] <= message.created_at]
                    sms_context[message.id] = prior[-1] if prior else None
        sms_link_profiles = (_app.db.session.scalars(select(
            SupporterProfile).order_by(SupporterProfile.name,
                                       SupporterProfile.id)).all()
                             if general_sms_history is not None else [])
        outbound_history = []
        if selected_contact is None:
            outbound_statement = select(_app.EmailMessage).order_by(
                _app.EmailMessage.created_at.desc(),
                _app.EmailMessage.id.desc()).limit(300)
            if not task_is_admin(user):
                assigned_family_ids = select(_app.FamilyAssignment.family_id).where(
                    _app.FamilyAssignment.staff_user_id == user.id)
                outbound_statement = outbound_statement.where(_app.or_(
                    _app.EmailMessage.family_id.in_(assigned_family_ids),
                    _app.EmailMessage.staff_user_id == user.id))
            outbound_history = _app.db.session.scalars(outbound_statement).all()
        pledge_delivery = {row.id: supporter_pledge_delivery(row) for row in contacts}
        return _app.render_template(
            'communications.html', title='Communications', contacts=contacts,
            selected_contact=selected_contact,
            history=history, latest=latest, due=due, email_replies=email_replies,
            applicant_history=applicant_history, applicant_replies=applicant_replies,
            inbox_history=inbox_history, inbox_messages=inbox_messages,
            general_sms_history=general_sms_history,
            general_sms_messages=general_sms_messages,
            sms_identities=sms_identities, sms_threads=sms_threads,
            sms_context=sms_context, sms_phone_key=sms_phone_key,
            sms_link_profiles=sms_link_profiles,
            outbound_history=outbound_history,
            callback_contact_ids=callback_contact_ids,
            overdue_contact_ids=overdue_contact_ids,
            pledge_delivery=pledge_delivery,
            pledge_frequencies=_app.PLEDGE_FREQUENCIES,
            ai_contact=ai_contact, ai_subject=ai_subject, ai_body=ai_body,
            external_email=external_email,
            sms_phone=sms_phone, sms_name=sms_name, sms_family=sms_family,
            supporter_query=supporter_query, supporter_page=supporter_page,
            supporters_have_next=supporters_have_next,
            now=now)

    def contact_mobile(contact):
        return contact.cell_phone or contact.phone or contact.home_phone or ''

    @app.route('/communications/case-broadcast', methods=['GET', 'POST'])
    def case_broadcast():
        user = task_user()
        if not user or user.role not in ('organization_admin', 'family_admin'):
            _app.abort(403)
        assigned_ids = select(_app.FamilyAssignment.family_id).where(
            _app.FamilyAssignment.staff_user_id == user.id)
        family_statement = select(_app.Family).order_by(_app.Family.name)
        if not task_is_admin(user):
            family_statement = family_statement.where(_app.Family.id.in_(assigned_ids))
        families = _app.db.session.scalars(family_statement).all()
        source = _app.request.form if _app.request.method == 'POST' else _app.request.args
        family_id = source.get('family_id', type=int)
        relationship = source.get('relationship', '')
        status_filter = source.get('status', '')
        channel = source.get('channel', 'email')
        if relationship and relationship not in set(_app.RELATIONSHIPS) | _app.LEGACY_RELATIONSHIPS:
            _app.abort(400)
        if status_filter and status_filter not in _app.CONTACT_STATUSES:
            _app.abort(400)
        if channel not in ('email', 'sms', 'whatsapp'):
            _app.abort(400)
        family = None
        recipients = []
        if family_id is not None:
            family = _app.db.get_or_404(_app.Family, family_id)
            if not task_is_admin(user) and family_id not in {row.id for row in families}:
                _app.abort(403)
            statement = select(_app.Contact).where(_app.Contact.family_id == family_id)
            if relationship:
                statement = statement.where(_app.Contact.relationship == relationship)
            if status_filter:
                statement = statement.where(_app.Contact.status == status_filter)
            contacts = _app.db.session.scalars(statement.order_by(_app.Contact.id)).all()
            seen = set()
            for contact in contacts:
                address = (contact.email or '').strip().lower() if channel == 'email' else contact_mobile(contact)
                if channel == 'email':
                    if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', address):
                        continue
                else:
                    try:
                        address = normalize_phone(address)
                    except ValueError:
                        continue
                if address not in seen:
                    seen.add(address)
                    recipients.append((contact, address))
        subject = source.get('subject', '').strip()[:300]
        body = source.get('body', '').strip()[:5000]
        if _app.request.method == 'POST' and source.get('confirm') == 'send':
            if not family or not recipients or not body or (channel == 'email' and not subject):
                _app.abort(400, 'Choose a case with eligible recipients and enter a message.')
            if channel != 'email' and len(body) > 1600:
                _app.abort(400, 'Message is too long.')
            if len(recipients) > 200:
                _app.abort(400, 'Narrow the category to 200 recipients or fewer.')
            sent = failed = 0
            for contact, address in recipients:
                if channel == 'email':
                    message = app.extensions['send_email'](
                        'case_broadcast', address, subject, body, family_id=family_id)
                    result = 'completed' if message.status == 'sent' else message.status
                    communication_row(contact, 'email', subject, body,
                                      status=result, email_message=message)
                else:
                    if app.config['TESTING'] or app.config['DEMO']:
                        provider_id, error, result = None, '', 'preview'
                    else:
                        provider_id, error = deliver_message(
                            app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
                            address, body, channel=channel,
                            sms_from=app.config['TWILIO_SMS_FROM'],
                            whatsapp_from=app.config['TWILIO_WHATSAPP_FROM'],
                            messaging_service_sid=twilio_service_sid())
                        result = 'failed' if error else 'completed'
                    communication_row(contact, channel, subject or channel.title(), body,
                                      status=result, provider_message_id=provider_id,
                                      delivery_error=error)
                sent += result == 'completed'
                failed += result == 'failed'
            add_audit(f'Case broadcast: {channel} to {family_id}; {sent} sent, {failed} failed')
            _app.db.session.commit()
            _app.flash(f'{sent} sent; {failed} failed; {len(recipients)-sent-failed} prepared.')
            return _app.redirect(_app.url_for('case_broadcast', family_id=family_id))
        return _app.render_template('case_broadcast.html', title='Case broadcast',
                                    families=families, family=family, recipients=recipients,
                                    relationships=_app.RELATIONSHIPS + sorted(_app.LEGACY_RELATIONSHIPS),
                                    relationship=relationship, status_filter=status_filter,
                                    channel=channel, subject=subject, body=body)

    def twilio_setting(key):
        row = _app.db.session.get(_app.OrganizationSetting, key)
        return row.value if row and isinstance(row.value, str) else ''

    def twilio_service_sid():
        return (app.config['TWILIO_MESSAGING_SERVICE_SID'] or
                twilio_setting('twilio_messaging_service_sid'))

    @app.get('/twilio-setup')
    def twilio_setup():
        if not task_is_admin(task_user()):
            _app.abort(403)
        account_sid = app.config['TWILIO_ACCOUNT_SID']
        auth_token = app.config['TWILIO_AUTH_TOKEN']
        overview, error = account_overview(account_sid, auth_token)
        test_delivery = None
        test_delivery_error = None
        test_message_sid = session.get('twilio_test_message_sid', '')
        if test_message_sid:
            test_delivery, test_delivery_error = message_status(
                account_sid, auth_token, test_message_sid)
        return _app.render_template(
            'twilio_setup.html', title='Twilio setup', overview=overview,
            error=error, account_sid_configured=bool(account_sid),
            auth_token_configured=bool(auth_token),
            sms_from=app.config['TWILIO_SMS_FROM'],
            messaging_service_sid=twilio_service_sid(),
            whatsapp_from=app.config['TWILIO_WHATSAPP_FROM'],
            inbound_webhook_url=(app.config.get('APP_BASE_URL', '').rstrip('/') +
                                 _app.url_for('twilio_incoming_message')),
            test_delivery=test_delivery,
            test_delivery_error=test_delivery_error)

    @app.post('/twilio-setup/messaging-service')
    def setup_twilio_messaging_service():
        if not task_is_admin(task_user()):
            _app.abort(403)
        phone_number_sid = _app.request.form.get('phone_number_sid', '').strip()
        if not re.fullmatch(r'PN[a-fA-F0-9]{32}', phone_number_sid):
            _app.abort(400, 'Choose a valid Twilio phone number.')
        overview, error = account_overview(
            app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'])
        owned = {row['sid'] for row in (overview or {}).get('numbers', [])}
        if error or phone_number_sid not in owned:
            _app.abort(400, error or 'That phone number does not belong to this Twilio account.')
        service_sid = find_messaging_service_for_number(
            app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
            overview.get('services', []), phone_number_sid)
        if service_sid:
            error = None
        else:
            service_sid, error = create_messaging_service(
                app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
                phone_number_sid)
        if error:
            _app.flash(f'Twilio setup failed: {error}', 'error')
        else:
            setting = _app.db.session.get(
                _app.OrganizationSetting, 'twilio_messaging_service_sid')
            if setting is None:
                setting = _app.OrganizationSetting(
                    key='twilio_messaging_service_sid', value=service_sid)
                _app.db.session.add(setting)
            else:
                setting.value = service_sid
            _app.db.session.commit()
            _app.flash('Yazory Messaging Service is connected.')
        return _app.redirect(_app.url_for('twilio_setup'))

    @app.post('/twilio-setup/test-message')
    def twilio_test_message():
        if not task_is_admin(task_user()):
            _app.abort(403)
        channel = _app.request.form.get('channel', '')
        if channel not in ('sms', 'whatsapp'):
            _app.abort(400, 'Choose SMS or WhatsApp.')
        recipient = _app.request.form.get('recipient_phone', '').strip()[:80]
        provider_id, error = deliver_message(
            app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
            recipient, 'Yazory Twilio connection test.', channel=channel,
            sms_from=app.config['TWILIO_SMS_FROM'],
            whatsapp_from=app.config['TWILIO_WHATSAPP_FROM'],
            messaging_service_sid=twilio_service_sid())
        if error:
            _app.flash(f'Test message failed: {error}', 'error')
        else:
            session['twilio_test_message_sid'] = provider_id
            _app.flash(
                f'Twilio accepted the test message. Reference: {provider_id}. '
                'Check the delivery status below.')
        return _app.redirect(_app.url_for('twilio_setup'))

    @app.post('/twilio/incoming-message')
    def twilio_incoming_message():
        signature = _app.request.headers.get('X-Twilio-Signature', '')
        public_base = app.config.get('APP_BASE_URL', '').rstrip('/')
        webhook_url = (public_base + _app.request.full_path.rstrip('?')
                       if public_base else _app.request.url)
        if not validate_webhook_signature(
                app.config['TWILIO_AUTH_TOKEN'], webhook_url,
                _app.request.form, signature):
            _app.abort(403)

        provider_id = (_app.request.form.get('MessageSid') or
                       _app.request.form.get('SmsSid') or '').strip()[:100]
        if provider_id and (
                _app.db.session.scalar(select(SupporterCommunication.id).where(
                    SupporterCommunication.provider_message_id == provider_id)) or
                _app.db.session.scalar(select(ApplicantMessage.id).where(
                    ApplicantMessage.provider_message_id == provider_id)) or
                _app.db.session.scalar(select(GeneralSmsMessage.id).where(
                    GeneralSmsMessage.provider_message_id == provider_id))):
            return Response('<Response></Response>', mimetype='application/xml')

        sender = _app.request.form.get('From', '').strip()[:80]
        body = _app.request.form.get('Body', '').strip()[:1600]
        channel = 'whatsapp' if sender.startswith('whatsapp:') else 'sms'
        family = family_for_inbound_phone(sender)
        contact = None if family else contact_for_inbound_phone(sender, channel)
        if family:
            _app.db.session.add(ApplicantMessage(
                family_id=family.id, direction='applicant', status='unread',
                body=body, provider_message_id=provider_id or None))
            _app.db.session.commit()
        elif contact:
            communication_row(
                contact, channel,
                'Incoming WhatsApp message' if channel == 'whatsapp'
                else 'Incoming text message',
                body, status='completed', provider_message_id=provider_id,
                direction='inbound')
            _app.db.session.commit()
        elif channel == 'sms':
            try:
                sender = normalize_phone(sender)
            except ValueError:
                sender = sender[:80]
            _app.db.session.add(GeneralSmsMessage(
                provider_message_id=provider_id or None, phone=sender,
                direction='inbound', body=body, status='unread'))
            _app.db.session.commit()
        return Response('<Response></Response>', mimetype='application/xml')

    @app.get('/communications')
    def communications():
        return render_communications()

    @app.post('/communications/external-email')
    def send_external_email():
        user = task_user()
        if user is None or user.role not in (
                'organization_admin', 'family_admin', 'fundraiser'):
            _app.abort(403)
        recipient = _app.request.form.get('recipient_email', '').strip().lower()[:254]
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', recipient):
            _app.abort(400, 'Enter a valid email address.')
        subject = _app.request.form.get('subject', '').strip()[:300]
        body = _app.request.form.get('body', '').strip()[:5000]
        if not subject or not body:
            _app.abort(400, 'Enter an email subject and message.')
        message = app.extensions['send_email'](
            'general_outbound', recipient, subject, body,
            staff_user_id=user.id, reply_to_public=True)
        _app.db.session.add(_app.Audit(
            actor=user.email,
            action=f'{"Sent" if message.status == "sent" else "Prepared"} general email to {recipient}'))
        _app.db.session.commit()
        if message.status == 'sent':
            _app.flash('Email sent.')
        elif message.status == 'preview':
            _app.flash(
                'Email was prepared but not sent because delivery is in preview mode.',
                'error')
        else:
            _app.flash('Email delivery failed. Check Email history.', 'error')
        return _app.redirect(_app.url_for('communications', _anchor='mailbox-sent'))

    @app.post('/communications/general-sms')
    def send_general_sms():
        user = task_user()
        if user is None or user.role != 'organization_admin':
            _app.abort(403)
        recipient = _app.request.form.get('recipient_phone', '').strip()[:80]
        contact_id = _app.request.form.get('contact_id', type=int)
        family_id = _app.request.form.get('family_id', type=int)
        family = None
        if family_id is not None:
            family = _app.db.get_or_404(_app.Family, family_id)
        if contact_id is not None:
            contact = _app.db.session.get(_app.Contact, contact_id)
            if contact is None:
                _app.abort(404)
            recipient = contact_mobile(contact)
            if not recipient:
                _app.abort(400, 'The selected person does not have a mobile number.')
        body = _app.request.form.get('body', '').strip()[:1600]
        if not body:
            _app.abort(400, 'Enter a message.')
        if not recipient:
            _app.abort(400, 'Choose a person or enter a mobile number.')
        try:
            recipient = normalize_phone(recipient)
        except ValueError as exc:
            _app.abort(400, str(exc))
        if app.config['TESTING'] or app.config['DEMO']:
            provider_id, error, status = None, '', 'preview'
        else:
            provider_id, error = deliver_message(
                app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
                recipient, body, channel='sms',
                sms_from=app.config['TWILIO_SMS_FROM'],
                messaging_service_sid=twilio_service_sid())
            status = 'failed' if error else 'completed'
        _app.db.session.add(GeneralSmsMessage(
            provider_message_id=provider_id or None, phone=recipient,
            direction='outbound', body=body, status=status,
            delivery_error=error or '', staff_user_id=user.id,
            family_id=family_id))
        _app.db.session.add(_app.Audit(
            actor=user.email,
            action=f'{"Sent" if status == "completed" else "Prepared" if status == "preview" else "Failed"} general SMS to {recipient}',
            family_id=family.id if family else None))
        _app.db.session.commit()
        if status == 'completed':
            _app.flash('SMS sent.')
        elif status == 'preview':
            _app.flash('SMS was prepared but not sent because delivery is in preview mode.', 'error')
        else:
            _app.flash('SMS delivery failed. Open Sent & history for the error.', 'error')
        return _app.redirect(_app.url_for('communications', _anchor='mailbox-sent'))

    @app.post('/resend/webhook')
    def resend_webhook():
        """Receive a verified supporter reply and add it to the case timeline."""
        raw_payload = _app.request.get_data(cache=False)
        signature_headers = {
            'svix-id': _app.request.headers.get('svix-id', ''),
            'svix-timestamp': _app.request.headers.get('svix-timestamp', ''),
            'svix-signature': _app.request.headers.get('svix-signature', ''),
        }
        try:
            event = verify_webhook(
                raw_payload, signature_headers,
                app.config.get('RESEND_WEBHOOK_SECRET', ''))
        except ValueError:
            _app.abort(400, 'Invalid webhook signature.')

        event_id = signature_headers['svix-id']
        event_type = event.get('type', '')
        data = event.get('data') or {}
        email_id = data.get('email_id', '') if event_type == 'email.received' else None
        if _app.db.session.get(ResendWebhookEvent, event_id) or (
                email_id and _app.db.session.scalar(select(ResendWebhookEvent).where(
                    ResendWebhookEvent.email_id == email_id))):
            return jsonify({'received': True, 'duplicate': True})
        if event_type != 'email.received':
            _app.db.session.add(ResendWebhookEvent(
                event_id=event_id, event_type=event_type or 'unknown'))
            _app.db.session.commit()
            return jsonify({'received': True, 'ignored': True})

        try:
            inbound = retrieve_received_email(app.config['RESEND_API_KEY'], email_id)
        except ValueError as exc:
            app.logger.warning('Could not retrieve Resend inbound email %s: %s', email_id, exc)
            # A non-2xx response makes Resend retry after a temporary API failure.
            return jsonify({'received': False, 'error': 'Email retrieval failed.'}), 503

        reply_domain = app.config.get('EMAIL_REPLY_DOMAIN', '')
        recipients = inbound.get('to') or data.get('to') or []
        if isinstance(recipients, str):
            recipients = [recipients]
        message = None
        if reply_domain:
            reply_pattern = re.compile(
                rf'^reply\+(\d+)-([0-9a-f]{{20}})@{re.escape(reply_domain)}$', re.I)
            for address in (item[1].lower() for item in getaddresses(recipients)):
                match = reply_pattern.fullmatch(address)
                if not match:
                    continue
                candidate = _app.db.session.get(_app.EmailMessage, int(match.group(1)))
                expected = hmac.new(
                    app.config['SECRET_KEY'].encode(), match.group(1).encode(),
                    hashlib.sha256).hexdigest()[:20]
                if candidate and hmac.compare_digest(expected, match.group(2).lower()):
                    message = candidate
                    break

        sender_values = inbound.get('from') or data.get('from') or ''
        parsed_senders = getaddresses(
            [sender_values] if isinstance(sender_values, str) else sender_values)
        sender_name, sender = next(((name, address.lower()) for name, address in parsed_senders
                                    if address), ('', ''))
        html_body = inbound.get('html') or ''
        body = (inbound.get('text') or html_to_text(html_body) or
                '[Email reply contained no readable text.]')[:50000]
        attachments = inbound.get('attachments') or []
        downloadable_attachments = [
            item for item in attachments
            if (item.get('content_disposition') or '').lower() != 'inline'
        ]
        if downloadable_attachments:
            names = ', '.join(str(item.get('filename') or 'attachment')[:255]
                              for item in downloadable_attachments[:20])
            body += f'\n\nAttachments: {names}'

        recipient_addresses = {
            address.lower() for _, address in getaddresses(recipients) if address
        }
        public_inbox = app.config.get('PUBLIC_INBOX_EMAIL', 'info@yaazory.org').lower()
        receiving_alias = (f'info@{reply_domain}' if reply_domain else '')
        if public_inbox in recipient_addresses or receiving_alias in recipient_addresses:
            _app.db.session.add(InboundInboxMessage(
                provider_message_id=email_id[:200],
                sender_name=(sender_name or '')[:200], sender_email=sender[:254],
                recipient=(public_inbox if public_inbox in recipient_addresses
                           else receiving_alias)[:254],
                subject=(inbound.get('subject') or data.get('subject') or
                         'Email to Yazory')[:300],
                body=body, html_body=html_body, status='unread'))
            _app.db.session.add(_app.Audit(
                actor=sender or 'Email sender', action='New message in Yazory inbox'))
            _app.db.session.add(ResendWebhookEvent(
                event_id=event_id, email_id=email_id, event_type=event_type))
            _app.db.session.commit()
            return jsonify({'received': True, 'matched': True,
                            'recipient': 'public_inbox'})

        # Staff emails to an applicant use the same signed Reply-To address.
        # Store a direct email reply in the applicant's portal conversation.
        if (message and message.kind == 'applicant_portal_message'
                and message.family_id and sender
                and sender == message.recipient.strip().lower()):
            family = _app.db.session.get(_app.Family, message.family_id)
            if family and sender == (family.email or '').strip().lower():
                _app.db.session.add(ApplicantMessage(
                    family_id=family.id, direction='applicant', status='unread',
                    body=body))
                _app.db.session.add(_app.Audit(
                    actor=f'Applicant: {family.email}',
                    action='Applicant replied by email', family_id=family.id))
                _app.db.session.add(ResendWebhookEvent(
                    event_id=event_id, email_id=email_id, event_type=event_type))
                _app.db.session.commit()
                return jsonify({'received': True, 'matched': True,
                                'recipient': 'applicant'})

        timeline = (_app.db.session.scalar(select(SupporterCommunication).where(
            SupporterCommunication.email_message_id == message.id).order_by(
                SupporterCommunication.id.desc())) if message else None)
        # The signed per-message address identifies the case; requiring the same
        # sender prevents a leaked reply address from adding mail to its timeline.
        if not timeline or not sender or sender != message.recipient.strip().lower():
            _app.db.session.add(ResendWebhookEvent(
                event_id=event_id, email_id=email_id, event_type=event_type))
            _app.db.session.commit()
            return jsonify({'received': True, 'unmatched': True})

        now = _app.datetime.now(_app.timezone.utc).replace(tzinfo=None)
        _app.db.session.add(SupporterCommunication(
            contact_id=timeline.contact_id, family_id=timeline.family_id,
            kind='email_reply', direction='inbound',
            subject=(inbound.get('subject') or data.get('subject') or 'Email reply')[:300],
            body=body, status='received', provider_message_id=email_id[:100],
            completed_at=now))
        _app.db.session.add(_app.Audit(
            actor=f'Supporter: {sender}', action='Supporter replied by email',
            family_id=timeline.family_id))
        task = _app.db.session.scalar(select(StaffTask).where(
            StaffTask.source_contact_id == timeline.contact_id))
        if task:
            task.status = 'To do'
            task.description = 'Supporter replied by email. Review the reply in Communications.'
            task.due_date = now.date()
            task.completed_at = None
        _app.db.session.add(ResendWebhookEvent(
            event_id=event_id, email_id=email_id, event_type=event_type))
        _app.db.session.commit()
        return jsonify({'received': True, 'matched': True})

    @app.post('/contacts/<int:contact_id>/communications/initial-email/draft')
    def draft_supporter_initial_email(contact_id):
        contact = communication_contact(contact_id)
        language = _app.session.get('language', 'en')
        try:
            body = draft_initial_email(
                os.environ.get('OPENAI_API_KEY', ''),
                os.environ.get('OPENAI_MODEL', 'gpt-5-mini'),
                language)
        except ValueError as exc:
            body = fallback_initial_email(language)
            _app.flash(
                'The AI service was unavailable, so an editable standard draft was opened.',
                'message')
        staff = task_user()
        body = body.replace('{supporter_name}', contact.name).replace(
            '{staff_name}', staff.name or 'the Yazory team')
        return render_communications(
            ai_contact=contact, ai_subject=initial_email_subject(language), ai_body=body)

    @app.post('/communications/replies/<int:reply_id>/handled')
    def handle_supporter_email_reply(reply_id):
        reply = _app.db.get_or_404(SupporterCommunication, reply_id)
        if reply.kind != 'email_reply' or reply.direction != 'inbound':
            _app.abort(404)
        communication_contact(reply.contact_id)
        reply.status = 'handled'
        reply.completed_at = _app.datetime.now(
            _app.timezone.utc).replace(tzinfo=None)
        _app.db.session.commit()
        _app.flash('Email reply marked as handled.')
        return _app.redirect(_app.url_for('communications', contact_id=reply.contact_id))

    @app.post('/communications/inbox/<int:message_id>/handled')
    def handle_inbound_inbox_message(message_id):
        user = task_user()
        if user is None or user.role != 'organization_admin':
            _app.abort(403)
        message = _app.db.get_or_404(InboundInboxMessage, message_id)
        message.status = 'handled'
        message.handled_at = _app.datetime.now(
            _app.timezone.utc).replace(tzinfo=None)
        message.handled_by = user.id
        _app.db.session.commit()
        _app.flash('Inbox message marked as handled.')
        return _app.redirect(_app.url_for('communications', _anchor='general-inbox'))

    @app.post('/communications/inbox/<int:message_id>/reply')
    def reply_to_inbound_inbox_message(message_id):
        user = task_user()
        if user is None or user.role != 'organization_admin':
            _app.abort(403)
        inbox_message = _app.db.get_or_404(InboundInboxMessage, message_id)
        recipient = (inbox_message.sender_email or '').strip().lower()
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', recipient):
            _app.abort(400, 'This message does not have a valid reply address.')
        subject = _app.request.form.get('subject', '').strip()[:300]
        body = _app.request.form.get('body', '').strip()[:5000]
        if not subject or not body:
            _app.abort(400, 'Enter an email subject and message.')

        sent = app.extensions['send_email'](
            'public_inbox_reply', recipient, subject, body,
            staff_user_id=user.id)
        if sent.status == 'sent':
            now = _app.datetime.now(_app.timezone.utc).replace(tzinfo=None)
            inbox_message.status = 'handled'
            inbox_message.handled_at = now
            inbox_message.handled_by = user.id
            _app.db.session.add(_app.Audit(
                actor=user.email,
                action=f'Replied to Yazory inbox message from {recipient}'))
            _app.flash('Reply sent and inbox message marked as handled.')
        elif sent.status == 'preview':
            _app.flash(
                'Reply was prepared but not sent because delivery is in preview mode.',
                'error')
        else:
            _app.flash('Reply delivery failed. Check Email history.', 'error')
        _app.db.session.commit()
        return _app.redirect(_app.url_for('communications', _anchor='general-inbox'))

    @app.post('/communications/general-sms/<int:message_id>/handled')
    def handle_general_sms(message_id):
        user = task_user()
        if user is None or user.role != 'organization_admin':
            _app.abort(403)
        message = _app.db.get_or_404(GeneralSmsMessage, message_id)
        if message.direction != 'inbound':
            _app.abort(404)
        message.status = 'handled'
        message.handled_at = _app.datetime.now(
            _app.timezone.utc).replace(tzinfo=None)
        _app.db.session.commit()
        _app.flash('SMS marked as handled.')
        return _app.redirect(_app.url_for('communications', _anchor='mailbox-inbox'))

    @app.post('/communications/general-sms/<int:message_id>/link')
    def link_general_sms_contact(message_id):
        user = task_user()
        if user is None or user.role != 'organization_admin':
            _app.abort(403)
        message = _app.db.get_or_404(GeneralSmsMessage, message_id)
        profile = _app.db.get_or_404(
            SupporterProfile, _app.request.form.get('profile_id', type=int))
        normalized = sms_phone_key(message.phone)
        if not normalized:
            _app.abort(400, 'This SMS number is not valid.')
        link = _app.db.session.get(SmsPhoneLink, normalized)
        if link is None:
            link = SmsPhoneLink(
                normalized_phone=normalized, profile_id=profile.id,
                created_by=user.id)
            _app.db.session.add(link)
        else:
            link.profile_id = profile.id
            link.created_by = user.id
            link.created_at = _app.datetime.now(
                _app.timezone.utc).replace(tzinfo=None)
        _app.db.session.commit()
        _app.flash('SMS number linked to existing person.')
        return _app.redirect(_app.url_for(
            'communications', _anchor=f'sms-{message.id}'))

    @app.post('/communications/general-sms/<int:message_id>/reply')
    def reply_to_general_sms(message_id):
        user = task_user()
        if user is None or user.role != 'organization_admin':
            _app.abort(403)
        message = _app.db.get_or_404(GeneralSmsMessage, message_id)
        if message.direction != 'inbound':
            _app.abort(404)
        body = _app.request.form.get('body', '').strip()[:1600]
        if not body:
            _app.abort(400, 'Enter a message.')
        if app.config['TESTING'] or app.config['DEMO']:
            provider_id, error, status = None, '', 'preview'
        else:
            provider_id, error = deliver_message(
                app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
                message.phone, body, channel='sms',
                sms_from=app.config['TWILIO_SMS_FROM'],
                messaging_service_sid=twilio_service_sid())
            status = 'failed' if error else 'completed'
        _app.db.session.add(GeneralSmsMessage(
            provider_message_id=provider_id or None, phone=message.phone,
            direction='outbound', body=body, status=status,
            delivery_error=error or '', staff_user_id=user.id,
            family_id=message.family_id))
        if status == 'completed':
            message.status = 'handled'
            message.handled_at = _app.datetime.now(
                _app.timezone.utc).replace(tzinfo=None)
            _app.flash('SMS reply sent and message marked as handled.')
        elif status == 'preview':
            _app.flash('SMS reply was prepared but not sent because delivery is in preview mode.', 'error')
        else:
            _app.flash('SMS reply failed. Open Sent & history for the error.', 'error')
        _app.db.session.commit()
        return _app.redirect(_app.url_for('communications', _anchor='mailbox-inbox'))

    def sms_notification_items(user, since):
        if user.role != 'organization_admin':
            return []
        messages = _app.db.session.scalars(select(GeneralSmsMessage).where(
            GeneralSmsMessage.direction == 'inbound',
            GeneralSmsMessage.status == 'unread').order_by(
            GeneralSmsMessage.created_at.desc(),
            GeneralSmsMessage.id.desc()).limit(100)).all()
        items = []
        for message in messages:
            identity = sms_person_identity(message.phone)
            display = identity['name'] if identity['known'] else 'Unknown contact'
            items.append({
                'id': f'sms-{message.id}', 'kind': 'Message',
                'action': f'Incoming SMS from {display}',
                'actor': message.phone, 'at': message.created_at,
                'url': _app.url_for('communications',
                                    _anchor=f'sms-{message.id}'),
                'direct_url': True,
                'profile_name': identity['name'] if identity['known'] else '',
                'profile_url': identity['url'] if identity['known'] else ''})
        supporter_replies = _app.db.session.scalars(select(
            SupporterCommunication).where(
            SupporterCommunication.kind == 'sms',
            SupporterCommunication.direction == 'inbound',
            SupporterCommunication.created_at > since).order_by(
            SupporterCommunication.created_at.desc(),
            SupporterCommunication.id.desc()).limit(100)).all()
        for message in supporter_replies:
            items.append({
                'id': f'supporter-sms-{message.id}', 'kind': 'Message',
                'action': f'Incoming SMS from {message.contact.name}',
                'actor': contact_mobile(message.contact) or message.contact.phone,
                'at': message.created_at,
                'url': _app.url_for('communications',
                                    contact_id=message.contact_id,
                                    _anchor='selected-supporter-history'),
                'direct_url': True, 'profile_name': message.contact.name,
                'profile_url': _app.url_for(
                    'supporter_detail', contact_id=message.contact_id)})
        return items

    app.extensions.setdefault('notification_item_providers', []).append(
        sms_notification_items)

    def validate_communication_task(contact_id):
        """Keep task form actions scoped to the task's own supporter."""
        task_id = _app.request.form.get('task_id', type=int)
        if task_id is not None:
            task = visible_task_or_403(task_id)
            if task.source_contact_id != contact_id:
                _app.abort(400, 'This task belongs to another supporter.')
            return task

    def communication_return(contact_id):
        task = validate_communication_task(contact_id)
        if task:
            return _app.redirect(_app.url_for('task_detail', task_id=task.id,
                                              _anchor='task-communications'))
        return _app.redirect(_app.url_for('communications', contact_id=contact_id))

    @app.post('/contacts/<int:contact_id>/communications/initial-email')
    def send_supporter_initial_email(contact_id):
        contact = communication_contact(contact_id)
        validate_communication_task(contact.id)
        recipient_email = _app.request.form.get('recipient_email', '').strip().lower()[:254]
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', recipient_email):
            _app.abort(400, 'Enter a valid email address.')
        subject = _app.request.form.get('subject', '').strip()[:300]
        body = _app.request.form.get('body', '').strip()[:5000]
        if not subject or not body:
            _app.abort(400, 'Enter an email subject and message.')
        message = app.extensions['send_email'](
            'supporter_initial_contact', recipient_email, subject, body,
            family_id=contact.family_id)
        communication_row(
            contact, 'initial_email', subject, body,
            status='completed' if message.status == 'sent' else message.status,
            email_message=message)
        update_supporter_person(contact, {'email': recipient_email})
        contact.status = 'To contact'
        task = _app.db.session.scalar(select(StaffTask).where(
            StaffTask.source_contact_id == contact.id))
        if task:
            task.status = 'Waiting'
            task.description = 'Waiting for supporter to reply to the initial email.'
        add_audit(f'{"Sent" if message.status == "sent" else "Prepared"} initial supporter email: {contact.name}')
        _app.db.session.commit()
        if message.status == 'sent':
            _app.flash('Initial email sent.')
        elif message.status == 'preview':
            _app.flash('Email was prepared but not sent because delivery is in preview mode.',
                       'error')
        else:
            _app.flash('Initial email delivery failed. Check Communications.', 'error')
        return communication_return(contact.id)

    @app.post('/contacts/<int:contact_id>/communications/callback')
    def schedule_supporter_callback(contact_id):
        contact = communication_contact(contact_id)
        validate_communication_task(contact.id)
        scheduled_for = parse_communication_time(
            _app.request.form.get('scheduled_for', ''))
        note = _app.request.form.get('note', '').strip()[:5000]
        previous = _app.db.session.scalars(select(SupporterCommunication).where(
            SupporterCommunication.contact_id == contact.id,
            SupporterCommunication.kind == 'callback',
            SupporterCommunication.status == 'scheduled')).all()
        for callback in previous:
            callback.status = 'cancelled'
        communication_row(contact, 'callback', 'Good time to call', note,
                          status='scheduled', scheduled_for=scheduled_for)
        task = _app.db.session.scalar(select(StaffTask).where(
            StaffTask.source_contact_id == contact.id))
        if task is None:
            assignee = automatic_task_assignee(contact)
            if assignee:
                task = StaffTask(
                    family_id=contact.family_id, source_contact_id=contact.id,
                    assigned_to=assignee.id, created_by=task_user().id,
                    title=f'Contact supporter: {contact.name}')
                _app.db.session.add(task)
        if task:
            task.status = 'Waiting'
            task.due_date = scheduled_for.date()
            task.description = note
            task.completed_at = None
            task.outcome = ''
        contact.status = 'To contact'
        add_audit(f'Scheduled supporter callback: {contact.name}')
        _app.db.session.commit()
        _app.flash('Callback saved.')
        return communication_return(contact.id)

    @app.post('/contacts/<int:contact_id>/communications/call')
    def complete_supporter_call(contact_id):
        contact = communication_contact(contact_id)
        validate_communication_task(contact.id)
        note = _app.request.form.get('note', '').strip()[:5000]
        if not note:
            _app.abort(400, 'Enter the result of the call before completing it.')
        outreach_status = _app.request.form.get('outreach_status', '').strip()
        if outreach_status not in TASK_OUTREACH_RESULTS:
            _app.abort(400, 'Choose the actual supporter outreach result.')
        communication_row(contact, 'phone_call', 'Phone call completed', note)
        now = _app.datetime.now(_app.timezone.utc).replace(tzinfo=None)
        pending_callbacks = _app.db.session.scalars(select(SupporterCommunication).where(
            SupporterCommunication.contact_id == contact.id,
            SupporterCommunication.kind == 'callback',
            SupporterCommunication.status == 'scheduled')).all()
        for callback in pending_callbacks:
            callback.status = 'completed'
            callback.completed_at = now
        linked = _app.db.session.scalars(select(_app.Contact).where(
            _app.Contact.supporter_key == contact.supporter_key)).all() \
            if contact.supporter_key else [contact]
        for row in linked:
            if row.status in ('To contact', 'No answer', 'Left a message',
                              'Call back', 'Contacted'):
                row.status = outreach_status
        task = _app.db.session.scalar(select(StaffTask).where(
            StaffTask.source_contact_id == contact.id))
        if task:
            task.status = 'Completed'
            task.completed_at = now
            task.outcome = note
        add_audit(f'Completed supporter call: {contact.name}')
        _app.db.session.commit()
        _app.flash('Phone call recorded.')
        return communication_return(contact.id)

    @app.post('/contacts/<int:contact_id>/communications/message/<channel>')
    def send_supporter_message(contact_id, channel):
        if channel not in ('sms', 'whatsapp'):
            _app.abort(404)
        contact = communication_contact(contact_id)
        validate_communication_task(contact.id)
        recipient = (_app.request.form.get('recipient_phone') or
                     contact.cell_phone or contact.phone or
                     contact.home_phone or '').strip()[:80]
        body = _app.request.form.get('body', '').strip()[:1600]
        if not body:
            _app.abort(400, 'Enter a message.')
        try:
            normalized = normalize_phone(recipient)
        except ValueError as exc:
            _app.abort(400, str(exc))

        if app.config['TESTING'] or app.config['DEMO']:
            provider_id, error, status = None, '', 'preview'
        else:
            provider_id, error = deliver_message(
                app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
                normalized, body, channel=channel,
                sms_from=app.config['TWILIO_SMS_FROM'],
                whatsapp_from=app.config['TWILIO_WHATSAPP_FROM'],
                messaging_service_sid=twilio_service_sid())
            status = 'failed' if error else 'completed'
        communication_row(
            contact, channel, 'WhatsApp message' if channel == 'whatsapp' else 'Text message',
            body, status=status, provider_message_id=provider_id,
            delivery_error=error)
        phone_values = {'cell_phone': normalized}
        if not contact.phone:
            phone_values['phone'] = normalized
        update_supporter_person(contact, phone_values)
        add_audit(
            f'{"Sent" if status == "completed" else "Prepared" if status == "preview" else "Failed"} '
            f'{channel} message: {contact.name}')
        _app.db.session.commit()
        if status == 'completed':
            _app.flash('Message sent.')
        elif status == 'preview':
            _app.flash('Message was prepared but not sent because delivery is in preview mode.', 'error')
        else:
            _app.flash('Message delivery failed. Open Communication history for the error.', 'error')
        return communication_return(contact.id)

    @app.post('/contacts/<int:contact_id>/communications/pledge')
    def send_supporter_pledge(contact_id):
        contact = communication_contact(contact_id)
        validate_communication_task(contact.id)
        if not contact.email:
            _app.abort(400, 'Enter the supporter email address before sending the pledge.')
        if contact.monthly_cents <= 0:
            _app.abort(400, 'Enter the pledge amount before sending it.')
        subject = 'Your Yazory pledge confirmation'
        delivery = supporter_pledge_delivery(contact)
        if not delivery['url']:
            _app.flash(
                'Add the public ABCharity campaign link before sending this pledge.',
                'error')
            return _app.redirect(_app.url_for(
                'charity_donations', family_id=delivery['family_id']))
        family_lines = '\n'.join(
            f'- {row.family.name}: ${row.monthly_cents / 100:,.2f} '
            f'{({"Weekly": "each week", "Monthly": "each month", "One time": "one time"}.get(row.pledge_frequency, row.pledge_frequency))}'
            for row in delivery['pledges'])
        destination = ('ABCharity campaign' if delivery['kind'] == 'ABCharity'
                       else 'secure Yazory payment page')
        body = (f'Dear {contact.name},\n\nThank you for your pledge through Yazory.\n\n'
                f'{family_lines}\n\n'
                f'Please use this {destination} to make your donation:\n'
                f'{delivery["url"]}\n\n'
                'A separate receipt will be emailed every time a payment is successfully received.\n\n'
                f'Your current pledge acknowledgment is available in your donor account after sign-in: '
                f'{_public_url("supporter_portal_pledge", contact_id=contact.id)}\n'
                'A pledge is not a donation receipt.\n\n'
                'If any detail is incorrect, please reply to this email before the next payment.')
        message = app.extensions['send_email'](
            'pledge_confirmation', contact.email, subject, body,
            family_id=contact.family_id)
        communication_row(contact, 'pledge_email', subject, body,
                          status='completed' if message.status == 'sent' else message.status,
                          email_message=message)
        linked = _app.db.session.scalars(select(_app.Contact).where(
            _app.Contact.supporter_key == contact.supporter_key)).all() \
            if contact.supporter_key else [contact]
        for row in linked:
            row.status = 'Pledged'
        add_audit(f'{"Sent" if message.status == "sent" else "Prepared"} supporter pledge via {delivery["kind"]}: {contact.name}')
        _app.db.session.commit()
        if message.status == 'sent':
            _app.flash('Pledge email sent.')
        elif message.status == 'preview':
            _app.flash('Email was prepared but not sent because delivery is in preview mode.',
                       'error')
        else:
            _app.flash('Pledge email delivery failed. Check Communications.', 'error')
        return communication_return(contact.id)

    @app.post('/contacts/<int:contact_id>/communications/pledge-details')
    def save_supporter_pledge_details(contact_id):
        contact = communication_contact(contact_id)
        validate_communication_task(contact.id)
        email = (_app.request.form.get('email') or '').strip()
        if email and (len(email) > 254 or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email)):
            _app.abort(400, 'Enter a valid supporter email address.')
        try:
            pledge_amount = Decimal(_app.request.form.get('monthly', '0'))
            if not pledge_amount.is_finite() or pledge_amount.as_tuple().exponent < -2:
                raise ValueError
            pledge_cents = int(pledge_amount * 100)
        except (InvalidOperation, TypeError, ValueError):
            _app.abort(400, 'Enter a valid pledge amount.')
        if pledge_cents < 0 or pledge_cents > 100000000:
            _app.abort(400, 'Enter a valid pledge amount.')
        frequency = (_app.request.form.get('pledge_frequency') or 'Monthly').strip()
        if frequency not in _app.PLEDGE_FREQUENCIES:
            _app.abort(400, 'Choose a valid donation frequency.')
        linked = _app.db.session.scalars(select(_app.Contact).where(
            _app.Contact.supporter_key == contact.supporter_key)).all() \
            if contact.supporter_key else [contact]
        for row in linked:
            row.email = email
            row.monthly_cents = pledge_cents
            row.pledge_frequency = frequency
        add_audit(f'Updated pledge details from communications: {contact.name}')
        _app.db.session.commit()
        _app.flash('Pledge details saved.')
        return communication_return(contact.id)

    def visible_task_or_403(task_id):
        task = _app.db.get_or_404(StaffTask, task_id)
        user = task_user()
        if not task_is_admin(user) and (user is None or task.assigned_to != user.id):
            _app.abort(403, 'You do not have permission to view this task.')
        return task

    def automatic_task_assignee(contact):
        """Choose the person already responsible for this supporter or case."""
        if app.extensions['workflows']['enforced']():
            link_model = app.extensions['workflows']['models']['SupporterLink']
            link = _app.db.session.get(link_model, contact.id)
            if link and link.assigned_to:
                assignee = _app.db.session.get(_app.StaffUser, link.assigned_to)
                if assignee and assignee.status == 'active':
                    return assignee
        assignee = _app.db.session.scalar(select(_app.StaffUser).join(
            _app.FamilyAssignment,
            _app.FamilyAssignment.staff_user_id == _app.StaffUser.id
        ).where(
            _app.FamilyAssignment.family_id == contact.family_id,
            _app.StaffUser.status == 'active',
            _app.StaffUser.role.in_(('fundraiser', 'family_admin')),
        ).order_by(
            (_app.StaffUser.role == 'fundraiser').desc(), _app.StaffUser.id
        ))
        if assignee:
            return assignee
        actor = task_user() if has_request_context() else None
        if actor and actor.status == 'active':
            return actor
        return _app.db.session.scalar(select(_app.StaffUser).where(
            _app.StaffUser.status == 'active',
            _app.StaffUser.role == 'organization_admin').order_by(_app.StaffUser.id))

    def sync_supporter_followup_task(contact):
        task = _app.db.session.scalar(select(StaffTask).where(
            StaffTask.source_contact_id == contact.id))
        if contact.status == 'To contact':
            assignee = automatic_task_assignee(contact)
            if assignee is None:
                return
            if task is None:
                task = StaffTask(
                    family_id=contact.family_id,
                    source_contact_id=contact.id,
                    assigned_to=assignee.id,
                    created_by=assignee.id,
                    title=f'Contact supporter: {contact.name}',
                    description='',
                    priority='Normal')
                _app.db.session.add(task)
                add_audit(f'Created automatic supporter follow-up: {contact.name}')
            else:
                task.assigned_to = assignee.id
            callback = scheduled_callback(contact.id)
            if callback:
                task.status = 'Waiting'
                task.due_date = callback.scheduled_for.date()
                task.completed_at = None
                task.outcome = ''
            elif task.status in ('Completed', 'Cancelled'):
                task.status = 'To do'
                task.completed_at = None
                task.outcome = ''
        elif task and task.status not in ('Completed', 'Cancelled'):
            task.status = ('Cancelled' if contact.status in ('Paused', 'Declined')
                           else 'Completed')
            task.completed_at = (_app.datetime.now(_app.timezone.utc).replace(tzinfo=None)
                                 if task.status == 'Completed' else None)
            task.outcome = (f'Outreach status updated to {contact.status}.'
                            if task.status == 'Completed' else '')

    def sync_contact_ids(contact_ids):
        for contact_id in contact_ids:
            contact = _app.db.session.get(_app.Contact, contact_id)
            if contact is not None:
                sync_supporter_followup_task(contact)
        _app.db.session.commit()

    app.extensions['sync_supporter_followup_task'] = sync_supporter_followup_task

    def backfill_missing_supporter_tasks():
        """Create only missing follow-up tasks for existing open supporters."""
        missing_ids = _app.db.session.scalars(
            select(_app.Contact.id).outerjoin(
                StaffTask, StaffTask.source_contact_id == _app.Contact.id
            ).where(
                _app.Contact.status == 'To contact',
                StaffTask.id.is_(None),
            ).order_by(_app.Contact.id)
        ).all()
        if missing_ids:
            sync_contact_ids(missing_ids)

    def repair_page_data():
        """Run legacy page-data repairs explicitly, never during a GET."""
        backfill_missing_supporter_tasks()
        candidates = _app.db.session.execute(select(
            _app.Contact, _app.EmailMessage.recipient
        ).join(
            SupporterCommunication,
            SupporterCommunication.contact_id == _app.Contact.id,
        ).join(
            _app.EmailMessage,
            _app.EmailMessage.id == SupporterCommunication.email_message_id,
        ).where(
            _app.Contact.email == '',
            _app.EmailMessage.status == 'sent',
            _app.EmailMessage.recipient != '',
        ).order_by(
            _app.Contact.id, SupporterCommunication.created_at.desc(),
            SupporterCommunication.id.desc(),
        )).all()
        repaired = set()
        for contact, recipient in candidates:
            address = recipient.strip().lower()
            if (contact.id not in repaired and not contact.email and
                    re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', address)):
                update_supporter_person(contact, {'email': address[:254]})
                repaired.add(contact.id)
        _app.db.session.commit()
        return len(repaired)

    app.extensions['repair_page_data'] = repair_page_data

    @app.cli.command('repair-page-data')
    def repair_page_data_command():
        """Backfill legacy task and supporter-email data outside requests."""
        repaired_emails = repair_page_data()
        print(f'Page data repaired; {repaired_emails} supporter emails restored.')

    def scheduled_callback(contact_id):
        if not contact_id:
            return None
        return _app.db.session.scalar(select(SupporterCommunication).where(
            SupporterCommunication.contact_id == contact_id,
            SupporterCommunication.kind == 'callback',
            SupporterCommunication.status == 'scheduled').order_by(
                SupporterCommunication.created_at.desc(),
                SupporterCommunication.id.desc()))

    def scheduled_callback_map(tasks):
        ids = {task.source_contact_id for task in tasks if task.source_contact_id}
        if not ids:
            return {}
        callbacks = _app.db.session.scalars(select(SupporterCommunication).where(
            SupporterCommunication.contact_id.in_(ids),
            SupporterCommunication.kind == 'callback',
            SupporterCommunication.status == 'scheduled').order_by(
                SupporterCommunication.created_at.desc(),
                SupporterCommunication.id.desc())).all()
        result = {}
        for callback in callbacks:
            result.setdefault(callback.contact_id, callback)
        return result

    app.extensions['scheduled_callback'] = scheduled_callback
    app.extensions['scheduled_callback_map'] = scheduled_callback_map

    def parsed_due_date():
        raw = _app.request.form.get('due_date', '').strip()
        if not raw:
            return None
        try:
            return _app.date.fromisoformat(raw)
        except ValueError:
            _app.abort(400, 'Enter a valid due date.')

    def active_assignee():
        assignee = _app.db.session.get(
            _app.StaffUser, _app.request.form.get('assigned_to', type=int))
        if assignee is None or assignee.status != 'active':
            _app.abort(400, 'Choose an active staff member or collector.')
        return assignee

    @app.route('/tasks', methods=['GET', 'POST'])
    def tasks():
        user = task_user()
        if user is None and not app.config['DEMO']:
            _app.abort(403)
        if _app.request.method == 'POST':
            if not task_is_admin(user):
                _app.abort(403, 'Only an organization administrator can assign a task.')
            title = _app.request.form.get('title', '').strip()[:240]
            if not title:
                _app.abort(400, 'Task title is required.')
            assignee = active_assignee()
            family_id = _app.request.form.get('family_id', type=int)
            if family_id and _app.db.session.get(_app.Family, family_id) is None:
                _app.abort(400, 'Choose a valid family case.')
            priority = _app.request.form.get('priority', 'Normal')
            if priority not in TASK_PRIORITIES:
                _app.abort(400, 'Choose a valid priority.')
            task = StaffTask(
                title=title,
                description=_app.request.form.get('description', '').strip()[:5000],
                assigned_to=assignee.id,
                created_by=user.id,
                family_id=family_id,
                priority=priority,
                due_date=parsed_due_date())
            _app.db.session.add(task)
            add_audit(f'Created task: {title} / assigned to {assignee.email}')
            _app.db.session.commit()
            _app.flash('Task created and assigned.')
            return _app.redirect(_app.url_for('task_detail', task_id=task.id))

        statement = select(StaffTask).options(
            joinedload(StaffTask.assignee), joinedload(StaffTask.family),
            joinedload(StaffTask.source_contact), selectinload(StaffTask.subtasks))
        if not task_is_admin(user):
            statement = statement.where(StaffTask.assigned_to == (user.id if user else -1))
        selected_status = _app.request.args.get('status', '').strip()
        if selected_status:
            if selected_status not in TASK_STATUSES:
                _app.abort(400)
            statement = statement.where(StaffTask.status == selected_status)

        selected_assignee_id = _app.request.args.get('assigned_to', type=int)
        if selected_assignee_id:
            statement = statement.where(StaffTask.assigned_to == selected_assignee_id)
        selected_family_id = _app.request.args.get('family_id', type=int)
        if selected_family_id:
            statement = statement.where(StaffTask.family_id == selected_family_id)
        selected_priority = _app.request.args.get('priority', '').strip()
        if selected_priority:
            if selected_priority not in TASK_PRIORITIES:
                _app.abort(400)
            statement = statement.where(StaffTask.priority == selected_priority)

        selected_due = _app.request.args.get('due', '').strip()
        if selected_due:
            if selected_due not in ('open', 'today', 'overdue'):
                _app.abort(400)
            from zoneinfo import ZoneInfo
            eastern_today = _app.datetime.now(ZoneInfo('America/New_York')).date()
            statement = statement.where(StaffTask.status.notin_(('Completed', 'Cancelled')))
            if selected_due == 'today':
                statement = statement.where(StaffTask.due_date == eastern_today)
            elif selected_due == 'overdue':
                statement = statement.where(StaffTask.due_date < eastern_today)

        status_rank = case(
            (StaffTask.status == 'In progress', 0),
            (StaffTask.status == 'To do', 1),
            (StaffTask.status == 'Waiting', 2),
            (StaffTask.status == 'Completed', 3),
            else_=4)
        priority_rank = case(
            (StaffTask.priority == 'Urgent', 0),
            (StaffTask.priority == 'High', 1),
            (StaffTask.priority == 'Normal', 2),
            else_=3)
        page = max(_app.request.args.get('page', 1, type=int), 1)
        page_size = 100
        rows = _app.db.session.scalars(statement.order_by(
            status_rank, priority_rank, StaffTask.due_date.is_(None),
            StaffTask.due_date, StaffTask.id.desc()).offset(
                (page - 1) * page_size).limit(page_size + 1)).all()
        has_next = len(rows) > page_size
        rows = rows[:page_size]
        staff = _app.db.session.scalars(select(_app.StaffUser).where(
            _app.StaffUser.status == 'active').order_by(_app.StaffUser.name, _app.StaffUser.email)).all()
        families = (_app.db.session.scalars(select(_app.Family).order_by(_app.Family.name)).all()
                    if task_is_admin(user) else [])
        return _app.render_template(
            'tasks.html', title='Tasks', tasks=rows, staff=staff, families=families,
            callbacks=scheduled_callback_map(rows),
            task_statuses=TASK_STATUSES, task_priorities=TASK_PRIORITIES,
            selected_status=selected_status, selected_assignee_id=selected_assignee_id,
            selected_family_id=selected_family_id, selected_priority=selected_priority,
            selected_due=selected_due, page=page, has_next=has_next,
            may_assign=task_is_admin(user), today=_app.date.today())

    @app.get('/tasks/<int:task_id>')
    def task_detail(task_id):
        task = visible_task_or_403(task_id)
        staff = _app.db.session.scalars(select(_app.StaffUser).where(
            _app.StaffUser.status == 'active').order_by(
                _app.StaffUser.name, _app.StaffUser.email)).all()
        recent_communications = None
        if task.source_contact and task_user().role in (
                'organization_admin', 'family_admin', 'fundraiser'):
            try:
                communication_contact(task.source_contact_id)
            except Forbidden:
                pass
            else:
                recent_communications = _app.db.session.scalars(
                    select(SupporterCommunication).where(
                        SupporterCommunication.contact_id == task.source_contact_id
                    ).order_by(SupporterCommunication.created_at.desc(),
                               SupporterCommunication.id.desc()).limit(5)
                ).all()
        return _app.render_template(
            'task_detail.html', title='Task details', task=task,
            recent_communications=recent_communications,
            callback=scheduled_callback(task.source_contact_id),
            task_statuses=TASK_STATUSES, task_priorities=TASK_PRIORITIES,
            outreach_results=TASK_OUTREACH_RESULTS,
            pledge_frequencies=_app.PLEDGE_FREQUENCIES,
            staff=staff, may_assign=task_is_admin(task_user()), today=_app.date.today())

    @app.post('/tasks/<int:task_id>/status')
    def update_task_status(task_id):
        task = visible_task_or_403(task_id)
        status = _app.request.form.get('status', '')
        if status not in TASK_STATUSES:
            _app.abort(400, 'Choose a valid task status.')
        outcome = _app.request.form.get('outcome', '').strip()
        outreach_status = _app.request.form.get('outreach_status', '').strip()
        if len(outcome) > 5000:
            _app.abort(400, 'Enter a status note of up to 5000 characters.')
        if status == 'Completed':
            if not outcome:
                _app.abort(400, 'Enter a completion verdict (up to 5000 characters).')
            if task.source_contact and outreach_status not in TASK_OUTREACH_RESULTS:
                _app.abort(400, 'Choose the actual supporter outreach result.')
        previous_status = task.status
        task.status = status
        if outcome:
            task.outcome = outcome
        task.completed_at = (_app.datetime.now(_app.timezone.utc).replace(tzinfo=None)
                             if status == 'Completed' else None)
        if task.source_contact:
            pending_callbacks = _app.db.session.scalars(select(SupporterCommunication).where(
                SupporterCommunication.contact_id == task.source_contact_id,
                SupporterCommunication.kind == 'callback',
                SupporterCommunication.status == 'scheduled')).all()
            if pending_callbacks and status in ('Completed', 'Cancelled'):
                for callback in pending_callbacks:
                    callback.status = 'completed' if status == 'Completed' else 'cancelled'
                    callback.completed_at = task.completed_at if status == 'Completed' else None
                if status == 'Completed' and previous_status != 'Completed':
                    communication_row(task.source_contact, 'phone_call',
                                      'Phone call completed',
                                      outcome)
            linked = [task.source_contact]
            if task.source_contact.supporter_key:
                linked = _app.db.session.scalars(select(_app.Contact).where(
                    _app.Contact.supporter_key == task.source_contact.supporter_key)).all()
            if status == 'Completed':
                for contact in linked:
                    if contact.status in ('To contact', 'No answer', 'Left a message',
                                          'Call back', 'Contacted'):
                        contact.status = outreach_status
            elif status == 'Cancelled':
                for contact in linked:
                    if contact.status == 'To contact':
                        contact.status = 'Paused'
        add_audit(f'Changed task #{task.id} status to {status}' +
                  (f'; note: {outcome}' if outcome else ''))
        _app.db.session.commit()
        _app.flash('Task status updated.')
        return _app.redirect(_app.url_for('task_detail', task_id=task.id))

    @app.post('/tasks/<int:task_id>/subtasks')
    def add_subtask(task_id):
        parent = visible_task_or_403(task_id)
        if parent.parent_id is not None:
            _app.abort(400, 'Subtasks cannot contain another level of subtasks.')
        user = task_user()
        title = _app.request.form.get('title', '').strip()[:240]
        if not title:
            _app.abort(400, 'Subtask title is required.')
        assignee = active_assignee() if task_is_admin(user) else user
        priority = _app.request.form.get('priority', parent.priority)
        if priority not in TASK_PRIORITIES:
            _app.abort(400, 'Choose a valid priority.')
        subtask = StaffTask(
            parent_id=parent.id, family_id=parent.family_id,
            assigned_to=assignee.id, created_by=user.id, title=title,
            description=_app.request.form.get('description', '').strip()[:5000],
            priority=priority, due_date=parsed_due_date())
        _app.db.session.add(subtask)
        add_audit(f'Added subtask to task #{parent.id}: {title}')
        _app.db.session.commit()
        _app.flash('Subtask added.')
        return _app.redirect(_app.url_for('task_detail', task_id=parent.id))

    def wrap_receipt_source(endpoint, send_manual_email=False):
        original = app.view_functions.get(endpoint)
        if original is None:
            return

        def wrapped(*args, **kwargs):
            response = original(*args, **kwargs)
            receipt_ids = tuple(getattr(_app.g, 'created_receipt_ids', ()) or ())
            new_receipts = [_app.db.session.get(_app.Receipt, receipt_id)
                            for receipt_id in receipt_ids]
            new_receipts = [receipt for receipt in new_receipts if receipt is not None]
            for receipt in new_receipts:
                if _app.db.session.scalar(select(SupporterCommunication.id).where(
                        SupporterCommunication.receipt_id == receipt.id)):
                    continue
                message = None
                if send_manual_email and receipt.contact.email:
                    subject = 'Your Yazory donation receipt'
                    body = (f'Thank you for your donation of '
                            f'${receipt.amount_cents / 100:,.2f}.\n\n'
                            f'Date received: {receipt.received_on.strftime("%m/%d/%Y")}\n'
                            f'Receipt reference: {receipt.reference or f"YZ-{receipt.id:06d}"}\n'
                            f'Download your receipt after donor sign-in: '
                            f'{_public_url("supporter_portal_receipt", receipt_id=receipt.id)}\n\n'
                            'Yazory is developed and operated by Synccos Inc.')
                    message = app.extensions['send_email'](
                        'donation_receipt', receipt.contact.email, subject, body,
                        family_id=receipt.family_id)
                row = communication_row(
                    receipt.contact, 'receipt_email', 'Donation receipt',
                    f'${receipt.amount_cents / 100:,.2f} · '
                    f'{receipt.reference or f"YZ-{receipt.id:06d}"}',
                    status=('failed' if message and message.status == 'failed'
                            else 'completed'), email_message=message)
                row.receipt_id = receipt.id
            if new_receipts:
                _app.db.session.commit()
            return response

        wrapped.__name__ = original.__name__
        app.view_functions[endpoint] = wrapped

    wrap_receipt_source('record_receipt', send_manual_email=True)
    wrap_receipt_source('stripe_webhook')

    @app.get('/supporters/<int:contact_id>/pledge-preview.pdf')
    def staff_pledge_preview(contact_id):
        from flask import send_file
        from receipt_pdf import build_pledge_acknowledgment_pdf
        contact = communication_contact(contact_id)
        if contact.status != 'Pledged' or contact.monthly_cents <= 0:
            _app.abort(404)
        return send_file(build_pledge_acknowledgment_pdf(contact), mimetype='application/pdf',
                         as_attachment=False)

    @app.get('/supporters/receipts/<int:receipt_id>/preview.pdf')
    def staff_receipt_preview(receipt_id):
        from flask import send_file
        from receipt_pdf import build_donor_receipt_pdf
        receipt = _app.db.get_or_404(_app.Receipt, receipt_id)
        communication_contact(receipt.contact_id)
        return send_file(build_donor_receipt_pdf(receipt), mimetype='application/pdf',
                         as_attachment=False)

    @app.get('/supporters/abcharity/<int:donation_id>/preview.pdf')
    def staff_abcharity_preview(donation_id):
        from flask import send_file
        from receipt_pdf import build_abcharity_payment_pdf
        donation = _app.db.get_or_404(_app.CharityDonation, donation_id)
        if not donation.donor.contact_id:
            _app.abort(404)
        contact = communication_contact(donation.donor.contact_id)
        return send_file(build_abcharity_payment_pdf(donation, contact),
                         mimetype='application/pdf', as_attachment=False)

    @app.post('/supporters/donations/<source>/<int:payment_id>/text')
    def staff_text_donation_record(source, payment_id):
        from supporter_portal import create_supporter_sms_signin_link
        if source == 'receipt':
            payment = _app.db.get_or_404(_app.Receipt, payment_id)
            contact = communication_contact(payment.contact_id)
            reference = payment.reference or f'YZ-{payment.id:06d}'
            amount_cents = payment.amount_cents
        elif source == 'abcharity':
            payment = _app.db.get_or_404(_app.CharityDonation, payment_id)
            if not payment.donor.contact_id:
                _app.abort(404)
            contact = communication_contact(payment.donor.contact_id)
            reference = f'ABCharity #{payment.external_id}'
            amount_cents = payment.amount_cents
        else:
            _app.abort(404)
        try:
            phone, link = create_supporter_sms_signin_link(app, contact)
        except ValueError as exc:
            _app.abort(400, str(exc))
        body = (f'Yazory: Your ${amount_cents / 100:,.2f} donation ({reference}) '
                f'is in your donor portal. Open your payment record here: {link} '
                'This one-time link expires in 30 minutes. Reply STOP to opt out.')
        if app.config['TESTING'] or app.config['DEMO']:
            provider_id, error, status = None, '', 'preview'
        else:
            provider_id, error = deliver_message(
                app.config['TWILIO_ACCOUNT_SID'], app.config['TWILIO_AUTH_TOKEN'],
                phone, body, channel='sms', sms_from=app.config['TWILIO_SMS_FROM'],
                messaging_service_sid=twilio_service_sid())
            status = 'failed' if error else 'completed'
        communication_row(contact, 'sms', 'Donation payment record', body,
                          status=status, provider_message_id=provider_id,
                          delivery_error=error)
        _app.db.session.commit()
        _app.flash('Text sent.' if status == 'completed' else
                   'Text prepared in preview mode.' if status == 'preview' else
                   'Text delivery failed. Check Communications.',
                   'success' if status == 'completed' else 'error')
        return _app.redirect(_app.url_for('supporter_detail', contact_id=contact.id))

    @app.post('/supporters/receipts/<int:receipt_id>/send')
    def staff_send_receipt(receipt_id):
        receipt = _app.db.get_or_404(_app.Receipt, receipt_id)
        contact = communication_contact(receipt.contact_id)
        if not contact.email:
            _app.abort(400, 'Add the donor email address before sending a receipt.')
        url = _public_url('supporter_portal_receipt', receipt_id=receipt.id)
        subject = 'Your Yazory donation receipt'
        body = (f'Dear {contact.name},\n\nHere is your receipt for the '
                f'${receipt.amount_cents / 100:,.2f} donation received on '
                f'{receipt.received_on.strftime("%m/%d/%Y")}.\n\n'
                f'Open it after signing in to your donor account: {url}')
        message = app.extensions['send_email']('donation_receipt', contact.email, subject,
                                                body, family_id=receipt.family_id)
        row = communication_row(contact, 'receipt_email', subject, body,
                                status='completed' if message.status == 'sent' else message.status,
                                email_message=message)
        row.receipt_id = receipt.id
        _app.db.session.commit()
        _app.flash('Receipt sent.' if message.status == 'sent' else
                   'Receipt delivery was not completed. Check Communications.',
                   'success' if message.status == 'sent' else 'error')
        return _app.redirect(_app.url_for('supporter_detail', contact_id=contact.id))

    @app.post('/supporters/<int:contact_id>/pledge-send')
    def staff_send_pledge_acknowledgment(contact_id):
        contact = communication_contact(contact_id)
        if contact.status != 'Pledged' or contact.monthly_cents <= 0:
            _app.abort(400, 'Record a pledge before sending its acknowledgment.')
        if not contact.email:
            _app.abort(400, 'Add the donor email address before sending a pledge acknowledgment.')
        url = _public_url('supporter_portal_pledge', contact_id=contact.id)
        subject = 'Your Yazory pledge acknowledgment'
        body = (f'Dear {contact.name},\n\nYour pledge of '
                f'${contact.monthly_cents / 100:,.2f} ({contact.pledge_frequency}) '
                f'for {contact.family.name} is recorded.\n\n'
                f'Open your acknowledgment after signing in: {url}\n\n'
                'A pledge is not a payment or a donation receipt.')
        message = app.extensions['send_email']('pledge_confirmation', contact.email,
                                                subject, body, family_id=contact.family_id)
        communication_row(contact, 'pledge_email', subject, body,
                          status='completed' if message.status == 'sent' else message.status,
                          email_message=message)
        _app.db.session.commit()
        _app.flash('Pledge acknowledgment sent.' if message.status == 'sent' else
                   'Pledge delivery was not completed. Check Communications.',
                   'success' if message.status == 'sent' else 'error')
        return _app.redirect(_app.url_for('supporter_detail', contact_id=contact.id))

    register_supporter_portal(app)
    register_applicant_portal(app)
    return register_native_payments(app)


if __name__ == '__main__':
    create_app().run(host='0.0.0.0', port=int(os.getenv('PORT', '5000')))
